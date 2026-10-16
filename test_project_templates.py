from __future__ import annotations

import importlib
import os
import sys
from pathlib import Path
from unittest.mock import Mock

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PyQt6.QtWidgets import QApplication, QFileDialog, QMessageBox

from CCTV_Report.cctv_report_app import (
    IMAGE_IMPORT_TEMPLATE_FILES,
    REPORT_TEMPLATE_FILES,
    MainWindow,
    example_template_path,
)

ROOT = Path(__file__).resolve().parent
TEMPLATE_DIR = ROOT / "contoh_template"


@pytest.fixture(scope="module")
def qt_app():
    return QApplication.instance() or QApplication([])


def _capture_modem_module(monkeypatch):
    capture_dir = ROOT / "Capture_Modem"
    if str(capture_dir) not in sys.path:
        sys.path.insert(0, str(capture_dir))

    import launch_guard

    monkeypatch.setenv(
        "NOC_LAUNCH_TOKEN", launch_guard.generate_launch_token()
    )
    return importlib.import_module("gui_app")


def test_capture_modem_sitelist_template_is_downloadable_and_parseable(
    qt_app, monkeypatch, tmp_path
) -> None:
    module = _capture_modem_module(monkeypatch)
    window = module.ModemCaptureGUI()
    output = tmp_path / "Capture Modem sitelist"
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        lambda *_args, **_kwargs: (str(output), "Excel Workbook (*.xlsx)"),
    )
    monkeypatch.setattr(QMessageBox, "information", Mock())
    try:
        assert window.btn_download_sitelist_template.text() == "Unduh Template .xlsx"
        window.btn_download_sitelist_template.click()

        downloaded = output.with_suffix(".xlsx")
        assert downloaded.is_file()
        workbook = load_workbook(downloaded)
        assert workbook["Site List"].values.__next__() == (
            "Site ID",
            "Site Name",
            "IP",
        )
        assert "Petunjuk" in workbook.sheetnames

        workbook["Site List"].append(["ABC00001", "Contoh Site", "192.0.2.10"])
        workbook.save(downloaded)
        rows = module.ip_lookup.parse_imported_sitelist_file(
            str(downloaded), mapping_list=[]
        )

        assert rows.to_dict("records") == [
            {
                "Site ID": "ABC00001",
                "Site Name": "Contoh Site",
                "IP": "192.0.2.10",
            }
        ]
    finally:
        window.close()


def test_cctv_template_menus_offer_report_and_image_import_packages(
    qt_app, monkeypatch
) -> None:
    window = MainWindow()
    calls: list[str] = []
    monkeypatch.setattr(
        window, "download_example_template", lambda filename: calls.append(filename)
    )
    try:
        expected_menus = (
            (window.download_report_template_button, REPORT_TEMPLATE_FILES),
            (window.download_image_template_button, IMAGE_IMPORT_TEMPLATE_FILES),
        )
        for button, templates in expected_menus:
            actions = button.menu().actions()
            assert [action.text() for action in actions] == list(templates)
            for action in actions:
                action.trigger()
            assert calls[-2:] == list(templates.values())
    finally:
        window.close()


def test_cctv_example_templates_download_for_both_packages(
    qt_app, monkeypatch, tmp_path
) -> None:
    window = MainWindow()
    messages = Mock()
    monkeypatch.setattr(QMessageBox, "information", messages)
    try:
        templates = [
            *REPORT_TEMPLATE_FILES.values(),
            *IMAGE_IMPORT_TEMPLATE_FILES.values(),
        ]
        for filename in templates:
            destination = tmp_path / filename
            monkeypatch.setattr(
                QFileDialog,
                "getSaveFileName",
                lambda *_args, target=destination, **_kwargs: (
                    str(target),
                    "Excel Workbook (*.xlsx)",
                ),
            )
            window.download_example_template(filename)

            source = example_template_path(filename)
            assert destination.is_file()
            assert destination.read_bytes() == source.read_bytes()
            workbook = load_workbook(destination, read_only=True, data_only=True)
            assert workbook.sheetnames
            assert workbook["Site List"].max_row > 1
        assert messages.call_count == len(templates)
    finally:
        window.close()


def test_example_templates_are_included_in_pyinstaller_specs() -> None:
    desktop_spec = (ROOT / "NOC Tools.spec").read_text(encoding="utf-8")
    modem_spec = (ROOT / "Capture_Modem.spec").read_text(encoding="utf-8")

    assert "EXAMPLE_TEMPLATES" in desktop_spec
    assert "contoh_template" in desktop_spec
    assert "Template_Sitelist_Capture_Modem.xlsx" in modem_spec
