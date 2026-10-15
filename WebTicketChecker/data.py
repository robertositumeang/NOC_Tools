from __future__ import annotations

import csv
from datetime import date, datetime
import io
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import xlrd


GAP_EXPORT_HEADERS = [
    "No",
    "Site ID",
    "Site Name",
    "Site ID OPSEL",
    "Province",
    "Regency/City",
    "Consortium",
    "Vendor MSS",
    "MSS Package",
    "Opsel",
    "Package",
    "Opsel Area Package",
    "Alarm ID",
    "Alarm Title",
    "Ticket ID",
    "Ticket Duration",
    "State Condition",
    "Status Planned",
    "Site Down Duration",
    "Date Time Site Down",
    "Down Duration Category",
    "Force Majeure Status",
    "Status",
    "Administrative State 1",
    "Administrative State 2",
    "Operation State",
    "Root Cause Analysis",
    "Root Cause Description",
]


def normalize_column_name(value: Any) -> str:
    return re.sub(r"[\s_\-/]+", "", str(value or "").strip()).casefold()


def get_column_value(row: dict[str, Any], candidates: Iterable[str]) -> str:
    normalized = {normalize_column_name(key): value for key, value in row.items()}
    for candidate in candidates:
        value = normalized.get(normalize_column_name(candidate))
        if value is not None and value != "":
            return str(value).strip()
    return ""


def normalize_site_id(value: Any) -> str:
    cleaned = str(value or "").strip()
    if re.fullmatch(r"\d+\.0+", cleaned):
        return cleaned.split(".", 1)[0]
    return cleaned


def read_incident_file(file_path: str | Path) -> list[dict[str, Any]]:
    path = Path(file_path)
    extension = path.suffix.casefold()
    if extension in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, data_only=True, read_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        if not rows:
            return []
        headers = [str(header).strip() if header is not None else "" for header in rows[0]]
        records: list[dict[str, Any]] = []
        for raw_row in rows[1:]:
            if not any(value not in (None, "") for value in raw_row):
                continue
            record: dict[str, Any] = {}
            for index, header in enumerate(headers):
                if not header:
                    continue
                value = raw_row[index] if index < len(raw_row) else ""
                if isinstance(value, (datetime, date)):
                    record[header] = value.isoformat(sep=" ")
                elif value is not None:
                    record[header] = str(value).strip()
                else:
                    record[header] = ""
            if any(record.values()):
                records.append(record)
        return records

    if extension == ".xls":
        book = xlrd.open_workbook(str(path))
        sheet = book.sheet_by_index(0)
        if sheet.nrows == 0:
            return []
        headers = [str(sheet.cell_value(0, col)).strip() for col in range(sheet.ncols)]
        records = []
        for row_idx in range(1, sheet.nrows):
            row_values = sheet.row_values(row_idx)
            if not any(val not in ("", None) for val in row_values):
                continue
            record = {}
            for col_idx, header in enumerate(headers):
                if not header:
                    continue
                val = row_values[col_idx]
                if isinstance(val, float) and val.is_integer():
                    record[header] = str(int(val))
                elif val is not None:
                    record[header] = str(val).strip()
                else:
                    record[header] = ""
            if any(record.values()):
                records.append(record)
        return records

    raise ValueError(f"File harus berformat Excel .xlsx, .xlsm, atau .xls. Ekstensi {path.suffix} tidak didukung.")


def compare_ticket_gaps(
    database_tickets: list[dict[str, Any]],
    incident_rows: list[dict[str, Any]],
    incident_type: str = "down",
) -> list[dict[str, Any]]:
    open_site_ids: set[str] = set()
    history_map: dict[str, list[dict[str, Any]]] = {}

    for row in database_tickets:
        site_id = normalize_site_id(
            row.get("siteId") or row.get("Site ID") or get_column_value(row, ["Site ID", "SiteID", "id"])
        )
        if not site_id:
            continue
        status = str(row.get("status") or row.get("Status TT") or row.get("Status") or "").strip().upper()
        if status == "OPEN":
            open_site_ids.add(site_id)
        history_map.setdefault(site_id, []).append(row)

    def _history_sort_key(item: dict[str, Any]) -> float:
        no_val = item.get("No") or item.get("no")
        try:
            if no_val is not None:
                return float(no_val)
        except (TypeError, ValueError):
            pass

        val = item.get("siteDownTimestamp") or item.get("siteDownTime") or item.get("dateQuery")
        if isinstance(val, (int, float)):
            return float(val)
        try:
            return float(val)
        except (TypeError, ValueError):
            pass
        return 0.0

    for history in history_map.values():
        history.sort(key=_history_sort_key, reverse=True)

    gaps: list[dict[str, Any]] = []
    for row in incident_rows:
        site_id = normalize_site_id(get_column_value(row, ["Site ID", "SiteID", "id"]))
        if not site_id:
            continue
        if site_id in open_site_ids:
            continue

        site_name = get_column_value(row, ["Site Name", "SiteName", "name"])
        province = get_column_value(row, ["Province", "Provinsi", "province"])
        regency = get_column_value(row, ["Regency/City", "Regency", "City", "Kab/Kota", "Kabupaten"])
        ticket_id = get_column_value(row, ["Ticket ID", "TicketID", "ticket id"])

        if incident_type == "down":
            duration = get_column_value(row, ["Site Down Duration", "Down Duration", "Duration"])
            dur_cat = get_column_value(row, ["Down Duration Category", "Duration Category"])
        else:
            duration = get_column_value(row, ["Site Unmonitor Duration", "Unmonitor Duration", "Duration"])
            dur_cat = get_column_value(row, ["Unmonitor Duration Category", "Duration Category"])

        history = history_map.get(site_id, [])
        if history:
            latest = history[0]
            db_status = str(latest.get("status") or latest.get("Status TT") or "-")
            db_ticket = str(latest.get("ticketId") or latest.get("Ticket ID") or "-")
        else:
            db_status = "Tidak ada di DB"
            db_ticket = "-"

        gaps.append(
            {
                "type": incident_type,
                "siteId": site_id,
                "siteName": site_name,
                "province": province,
                "regency": regency,
                "ticketId": ticket_id,
                "duration": duration,
                "durCat": dur_cat,
                "dbStatus": db_status,
                "dbTicket": db_ticket,
                "history": history,
                "_history": history,
                "_raw": row,
                "raw": row,
            }
        )
    return gaps


def audit_both_incident_files(
    database_tickets: list[dict[str, Any]],
    down_rows: list[dict[str, Any]],
    unmon_rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    down_gaps = compare_ticket_gaps(database_tickets, down_rows, "down")
    unmon_gaps = compare_ticket_gaps(database_tickets, unmon_rows, "unmon")
    all_gaps = down_gaps + unmon_gaps

    open_site_ids = {
        normalize_site_id(row.get("siteId") or row.get("Site ID") or "")
        for row in database_tickets
        if str(row.get("status") or row.get("Status TT") or row.get("Status") or "").strip().upper() == "OPEN"
    }

    down_ok = sum(
        1
        for row in down_rows
        if normalize_site_id(get_column_value(row, ["Site ID", "SiteID", "id"])) in open_site_ids
    )
    unmon_ok = sum(
        1
        for row in unmon_rows
        if normalize_site_id(get_column_value(row, ["Site ID", "SiteID", "id"])) in open_site_ids
    )

    summary = {
        "open": down_ok + unmon_ok,
        "gaps": len(all_gaps),
        "total": len(down_rows) + len(unmon_rows),
        "down": len(down_gaps),
        "unmonitor": len(unmon_gaps),
        "down_total": len(down_rows),
        "unmon_total": len(unmon_rows),
        "down_gaps": len(down_gaps),
        "unmon_gaps": len(unmon_gaps),
        "total_gaps": len(all_gaps),
        "total_checked": len(down_rows) + len(unmon_rows),
        "ok_in_db": down_ok + unmon_ok,
    }
    return all_gaps, summary


def write_gap_workbook(records: list[dict[str, Any]], destination: str | Path) -> Path:
    if not records:
        raise ValueError("Tidak ada gap tiket untuk diekspor.")

    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Belum Open TT"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="102238", end_color="102238", fill_type="solid"
    )
    center_align = Alignment(horizontal="center", vertical="center")

    sheet.append(GAP_EXPORT_HEADERS)
    for col_index in range(1, len(GAP_EXPORT_HEADERS) + 1):
        cell = sheet.cell(row=1, column=col_index)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

def extract_row_values_for_template(item: dict[str, Any], record_index: int) -> list[Any]:
    raw = item.get("_raw") if isinstance(item, dict) and "_raw" in item else item
    if not isinstance(raw, dict):
        raw = {}

    lookup = {normalize_column_name(k): v for k, v in raw.items()}

    def get_val(keys: list[str], default: str = "") -> str:
        for k in keys:
            norm = normalize_column_name(k)
            if norm in lookup and lookup[norm] is not None and str(lookup[norm]).strip() != "":
                return str(lookup[norm]).strip()
        return default

    row_values: list[Any] = []
    for header in GAP_EXPORT_HEADERS:
        if header == "No":
            row_values.append(record_index)
        elif header == "Site ID":
            row_values.append(get_val(["Site ID", "SiteID", "id"], str(item.get("siteId", ""))))
        elif header == "Site Name":
            row_values.append(get_val(["Site Name", "SiteName", "Nama Site"], str(item.get("siteName", ""))))
        elif header == "Site ID OPSEL":
            row_values.append(get_val(["Site ID OPSEL", "Site ID Opsel", "Opsel Site ID", "SiteID OPSEL"]))
        elif header == "Province":
            row_values.append(get_val(["Province", "Provinsi"], str(item.get("province", ""))))
        elif header == "Regency/City":
            row_values.append(get_val(["Regency/City", "Regency", "City", "Kabupaten/Kota", "Kabupaten", "Kota"], str(item.get("regency", ""))))
        elif header == "Consortium":
            row_values.append(get_val(["Consortium", "Konsorsium"]))
        elif header == "Vendor MSS":
            row_values.append(get_val(["Vendor MSS", "VendorMSS", "Vendor"]))
        elif header == "MSS Package":
            row_values.append(get_val(["MSS Package", "MSSPackage"]))
        elif header == "Opsel":
            row_values.append(get_val(["Opsel", "Operator"]))
        elif header == "Package":
            row_values.append(get_val(["Package", "Paket"], str(item.get("package", ""))))
        elif header == "Opsel Area Package":
            row_values.append(get_val(["Opsel Area Package"]))
        elif header == "Alarm ID":
            row_values.append(get_val(["Alarm ID", "AlarmID"]))
        elif header == "Alarm Title":
            row_values.append(get_val(["Alarm Title", "AlarmTitle", "Alarm"]))
        elif header == "Ticket ID":
            val = get_val(["Ticket ID", "TicketID"], str(item.get("ticketId", "-")))
            row_values.append(val if val else "-")
        elif header == "Ticket Duration":
            row_values.append(get_val(["Ticket Duration", "TicketDuration"], "-"))
        elif header == "State Condition":
            val = get_val(["State Condition", "StateCondition", "Condition"], str(item.get("type", "down")))
            row_values.append(val)
        elif header == "Status Planned":
            row_values.append(get_val(["Status Planned", "StatusPlanned", "Planned"], "Unplanned"))
        elif header == "Site Down Duration":
            val = get_val([
                "Site Down Duration",
                "Site Unmonitor Duration",
                "Down Duration",
                "Unmonitor Duration",
                "Duration",
                "Durasi",
            ], str(item.get("duration", "")))
            row_values.append(val)
        elif header == "Date Time Site Down":
            val = get_val([
                "Date Time Site Down",
                "Date Time Site Unmonitor",
                "Date Time Down",
                "Date Time Unmonitor",
                "Down Time",
                "Waktu Down",
            ])
            row_values.append(val)
        elif header == "Down Duration Category":
            val = get_val([
                "Down Duration Category",
                "Unmonitor Duration Category",
                "Duration Category",
                "Kategori Durasi",
            ], str(item.get("durCat", "")))
            row_values.append(val)
        elif header == "Force Majeure Status":
            row_values.append(get_val(["Force Majeure Status", "Force Majeure"], "-"))
        elif header == "Status":
            row_values.append(get_val(["Status"], "-"))
        elif header == "Administrative State 1":
            row_values.append(get_val(["Administrative State 1", "Admin State 1"], "-"))
        elif header == "Administrative State 2":
            row_values.append(get_val(["Administrative State 2", "Admin State 2"], "-"))
        elif header == "Operation State":
            row_values.append(get_val(["Operation State", "Op State"], "-"))
        elif header == "Root Cause Analysis":
            row_values.append(get_val(["Root Cause Analysis", "RCA"], "-"))
        elif header == "Root Cause Description":
            row_values.append(get_val(["Root Cause Description", "RCA Description"], "-"))
        else:
            row_values.append(get_val([header]))

    return row_values


def write_gap_workbook(records: list[dict[str, Any]], destination: str | Path) -> Path:
    if not records:
        raise ValueError("Tidak ada gap tiket untuk diekspor.")

    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Belum Open TT"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="102238", end_color="102238", fill_type="solid"
    )
    center_align = Alignment(horizontal="center", vertical="center")

    sheet.append(GAP_EXPORT_HEADERS)
    for col_index in range(1, len(GAP_EXPORT_HEADERS) + 1):
        cell = sheet.cell(row=1, column=col_index)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for record_index, item in enumerate(records, start=1):
        row_values = extract_row_values_for_template(item, record_index)
        sheet.append(row_values)

    for col in sheet.columns:
        first_val = str(col[0].value or "")
        max_len = max(len(str(c.value or "")) for c in col)
        col_letter = col[0].column_letter
        sheet.column_dimensions[col_letter].width = max(max_len + 3, len(first_val) + 3, 12)

    workbook.save(path)
    return path


def generate_gap_csv_text(records: list[dict[str, Any]]) -> str:
    if not records:
        raise ValueError("Tidak ada gap tiket untuk diekspor.")

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(GAP_EXPORT_HEADERS)

    for index, item in enumerate(records, start=1):
        row_values = extract_row_values_for_template(item, index)
        writer.writerow(row_values)

    return output.getvalue()


def write_gap_csv(records: list[dict[str, Any]], destination: str | Path) -> Path:
    csv_text = generate_gap_csv_text(records)
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(csv_text, encoding="utf-8-sig")
    return path


# ── Google Sheets Write Helpers ─────────────────────────────────────────────

# Kolom yang HARUS diisi ke Google Sheets (Database TT)
# Sesuai mapping: C=3, M=13, N=14, O=15, P=16, Q=17, S=19, X=24, Y=25, Z=26, AA=27
SHEET_WRITE_COLUMNS = [
    "Site ID",            # Col C (3)
    "Alarm ID",           # Col M (13)
    "Alarm Title",        # Col N (14)
    "Ticket ID",          # Col O (15)
    "Ticket Duration",    # Col P (16)
    "State Condition",    # Col Q (17)
    "Date Time Site Down",# Col S (19) — format DD/MM/YYYY hh:mm
    "Status",             # Col X (24) — BAPL / Non BAPL
    "Administrative State 1",  # Col Y (25)
    "Administrative State 2",  # Col Z (26)
    "Operation State",    # Col AA (27)
]

# Pemetaan kolom Database TT (1-indexed, sesuai spreadsheet)
SHEET_COLUMN_INDICES = {
    "Site ID": 3,
    "Alarm ID": 13,
    "Alarm Title": 14,
    "Ticket ID": 15,
    "Ticket Duration": 16,
    "State Condition": 17,
    "Date Time Site Down": 19,
    "Status": 24,
    "Administrative State 1": 25,
    "Administrative State 2": 26,
    "Operation State": 27,
}


def format_datetime_dd_mm_yyyy_hh_mm(value: Any) -> str:
    """
    Konversi berbagai format datetime ke DD/MM/YYYY hh:mm.
    Mendukung: isoformat, M/D/YYYY H:MM:SS, dan string tanggal umum.
    """
    if not value:
        return ""
    s = str(value).strip()
    if not s or s in ("-", "—", "None", "nan"):
        return ""

    # Coba parse berbagai format
    formats = [
        "%m/%d/%Y %H:%M:%S",   # 10/9/2026 10:30:00
        "%m/%d/%Y %H:%M",      # 10/9/2026 10:30
        "%Y-%m-%d %H:%M:%S",   # 2026-10-09 10:30:00
        "%Y-%m-%dT%H:%M:%S",   # 2026-10-09T10:30:00
        "%Y-%m-%d %H:%M",      # 2026-10-09 10:30
        "%d/%m/%Y %H:%M:%S",   # 09/10/2026 10:30:00
        "%d/%m/%Y %H:%M",      # 09/10/2026 10:30
        "%d-%m-%Y %H:%M:%S",   # 09-10-2026 10:30:00
        "%d-%m-%Y %H:%M",      # 09-10-2026 10:30
        "%m/%d/%Y",             # 10/9/2026
        "%Y-%m-%d",             # 2026-10-09
        "%d/%m/%Y",             # 09/10/2026
    ]
    for fmt in formats:
        try:
            dt = datetime.strptime(s, fmt)
            return dt.strftime("%d/%m/%Y %H:%M")
        except ValueError:
            continue

    # Kembalikan string asli jika tidak bisa di-parse
    return s


def extract_sheet_write_payload(
    items: list[dict[str, Any]],
    indices: list[int],
) -> list[dict[str, Any]]:
    """
    Ekstrak payload untuk ditulis ke Google Sheets Database TT.
    Hanya kolom yang diizinkan yang diisi.

    Args:
        items: List gap items dari _state['gaps']
        indices: Indeks baris yang dipilih (0-based, terhadap list items)

    Returns:
        List dict, setiap dict mewakili satu baris dengan key = nama kolom,
        value = nilai yang akan diisi ke Google Sheets.
    """
    result: list[dict[str, Any]] = []
    for idx in indices:
        if idx < 0 or idx >= len(items):
            continue
        item = items[idx]
        raw = item.get("_raw") or item.get("raw") or {}
        if not isinstance(raw, dict):
            raw = {}
        lookup = {normalize_column_name(k): v for k, v in raw.items()}

        def gv(keys: list[str], fallback: str = "") -> str:
            for k in keys:
                norm = normalize_column_name(k)
                if norm in lookup and lookup[norm] is not None:
                    v = str(lookup[norm]).strip()
                    if v and v not in ("-", "—", "None", "nan"):
                        return v
            return fallback

        site_id = gv(["Site ID", "SiteID", "id"], str(item.get("siteId", "")))
        alarm_id = gv(["Alarm ID", "AlarmID"])
        alarm_title = gv(["Alarm Title", "AlarmTitle", "Alarm"])
        ticket_id = gv(["Ticket ID", "TicketID"], str(item.get("ticketId", "")))
        ticket_dur = gv(["Ticket Duration", "TicketDuration"])
        state_cond = gv(["State Condition", "StateCondition", "Condition"],
                        str(item.get("type", "")))
        raw_dt = gv([
            "Date Time Site Down", "Date Time Site Unmonitor",
            "Date Time Down", "Date Time Unmonitor",
            "Down Time", "Waktu Down",
        ])
        date_down = format_datetime_dd_mm_yyyy_hh_mm(raw_dt)
        status = gv(["Status"])
        adm_state1 = gv(["Administrative State 1", "Admin State 1"])
        adm_state2 = gv(["Administrative State 2", "Admin State 2"])
        op_state = gv(["Operation State", "Op State"])

        row: dict[str, Any] = {
            "Site ID": site_id,
            "Alarm ID": alarm_id,
            "Alarm Title": alarm_title,
            "Ticket ID": ticket_id,
            "Ticket Duration": ticket_dur,
            "State Condition": state_cond,
            "Date Time Site Down": date_down,
            "Status": status,
            "Administrative State 1": adm_state1,
            "Administrative State 2": adm_state2,
            "Operation State": op_state,
            # Metadata (tidak ditulis ke sheet, hanya untuk UI)
            "_idx": idx,
            "_siteName": str(item.get("siteName", "")),
        }
        result.append(row)
    return result


def generate_tsv_paste_data(
    items: list[dict[str, Any]],
    indices: list[int],
) -> str:
    """
    Generate teks TSV (Tab Separated Values) siap paste ke Google Sheets.
    Kolom sesuai urutan SHEET_WRITE_COLUMNS.
    Baris pertama adalah header, diikuti data.
    """
    payload = extract_sheet_write_payload(items, indices)
    if not payload:
        return ""

    lines = ["\t".join(SHEET_WRITE_COLUMNS)]
    for row in payload:
        values = [str(row.get(col, "") or "") for col in SHEET_WRITE_COLUMNS]
        lines.append("\t".join(values))
    return "\n".join(lines)
