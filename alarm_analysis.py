from __future__ import annotations

import csv
from datetime import date, datetime, timedelta
from io import BytesIO, StringIO
from itertools import chain
import re
import zipfile
from typing import Any

import xlrd
from openpyxl import load_workbook


def _normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = str(value).strip()
    return "" if result.casefold() in {"", "none", "null", "-", "n/a"} else result


def _get(row: dict[str, Any], *names: str) -> Any:
    values = {_normalise_header(key): value for key, value in row.items()}
    for name in names:
        key = _normalise_header(name)
        if key in values:
            return values[key]
    return None


def _rows_from_grid(grid: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    if not grid:
        raise ValueError("Workbook has no header row.")
    headers = [_text(value) or f"Column {index + 1}" for index, value in enumerate(grid[0])]
    rows = []
    for values in grid[1:]:
        if not any(_text(value) for value in values):
            continue
        rows.append(dict(zip(headers, values)))
    if not rows:
        raise ValueError("Workbook has no data rows.")
    return rows


def format_duration(occurrence: datetime | None, clear_time: datetime | None) -> str:
    """Format the down duration into a human-readable string."""
    if occurrence is None:
        return "-"
    if clear_time is None:
        return "Active (Masih Down)"
    if clear_time < occurrence:
        return "-"
    total_seconds = int((clear_time - occurrence).total_seconds())
    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60
    seconds = total_seconds % 60
    if hours > 0:
        return f"{hours}h {minutes}m {seconds}s"
    if minutes > 0:
        return f"{minutes}m {seconds}s"
    return f"{seconds}s"


def detect_alarm_vendor(rows: list[dict[str, Any]]) -> str | None:
    """Identify an NMS export from alarm table fields and values."""
    if not rows:
        return None
    headers = {_normalise_header(header) for header in rows[0]}
    if {"faultsource", "generationtime", "alarmname"} <= headers or {"neid", "generationtime"} <= headers:
        return "Fiberhome"
    if {"enodebid", "occurredonnt"} <= headers or {"enodebid", "name"} <= headers or {"enodebname", "occurredonnt"} <= headers:
        return "Huawei"
    if {"alarmsource", "occurredonnt", "name", "netype"} <= headers:
        return "Huawei"
    if {"me", "alarmcodename", "occurrencetime"} <= headers or {"meid", "alarmcodename"} <= headers or {"alarmcodename", "occurrencetime"} <= headers:
        return "ZTE"
    if {"alarmsource", "occurredonnt", "name"} <= headers:
        alarm_names = (
            _text(_get(row, "Name", "Alarm Name")).casefold()
            for row in rows[:10]
        )
        if any(re.search(r"\bme\b", name) for name in alarm_names):
            return "ZTE"
        return "Huawei"
    return None


def _is_table_header(row: tuple[Any, ...]) -> bool:
    headers = {_normalise_header(value) for value in row if _text(value)}
    if {"siteid", "sitename"} <= headers:
        return True
    alarm_header_signatures = (
        {"faultsource", "generationtime", "alarmname"},
        {"me", "alarmcodename", "occurrencetime"},
        {"alarmcodename", "occurrencetime"},
        {"enodebid", "occurredonnt"},
        {"enodebid", "name"},
        {"enodebname", "occurredonnt"},
        {"alarmsource", "occurredonnt", "name"},
    )
    return any(signature <= headers for signature in alarm_header_signatures)


def _read_xlsx(payload: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    sheets_to_try = [workbook.active] if workbook.active else []
    for name in workbook.sheetnames:
        s = workbook[name]
        if s not in sheets_to_try:
            sheets_to_try.append(s)

    best_rows: list[dict[str, Any]] = []
    for sheet in sheets_to_try:
        if hasattr(sheet, "reset_dimensions"):
            # ZTE NMS exports specify a bogus <dimension ref="A1" /> in XML,
            # which tricks openpyxl read_only mode into stopping at cell A1.
            # reset_dimensions forces openpyxl to read all rows until EOF.
            sheet.reset_dimensions()
        rows_iter = iter(sheet.iter_rows(values_only=True))
        scanned_rows: list[tuple[Any, ...]] = []
        header_row: tuple[Any, ...] | None = None
        for _ in range(100):
            try:
                row = next(rows_iter)
            except StopIteration:
                break
            scanned_rows.append(row)
            if _is_table_header(row):
                header_row = row
                break
        if not scanned_rows:
            continue
        if header_row is None:
            header_row = scanned_rows[0]
            data_rows = chain(scanned_rows[1:], rows_iter)
        else:
            data_rows = rows_iter
        headers = [
            _text(value) or f"Column {index + 1}"
            for index, value in enumerate(header_row)
        ]
        sheet_rows = []
        for values in data_rows:
            if not any(_text(value) for value in values):
                continue
            sheet_rows.append(dict(zip(headers, values)))
        if sheet_rows:
            best_rows = sheet_rows
            break

    workbook.close()
    if not best_rows:
        raise ValueError("Workbook has no data rows.")
    return best_rows


def _read_xls(payload: bytes) -> list[dict[str, Any]]:
    workbook = xlrd.open_workbook(file_contents=payload)
    sheet = workbook.sheet_by_index(0)
    grid = []
    for row_index in range(sheet.nrows):
        values = []
        for column_index in range(sheet.ncols):
            cell = sheet.cell(row_index, column_index)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = xlrd.xldate_as_datetime(value, workbook.datemode)
            values.append(value)
        grid.append(tuple(values))
    return _rows_from_grid(grid)


def _read_csv(payload: bytes, filename: str) -> list[dict[str, Any]]:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = payload.decode("cp1252")
    grid = list(csv.reader(StringIO(text)))
    required_headers = (
        {"me", "alarmcodename", "occurrencetime"},
        {"alarmcodename", "occurrencetime"},
        {"alarmsource", "occurredonnt", "name"},
        {"enodebid", "occurredonnt"},
        {"occurredonnt", "name"},
        {"faultsource", "generationtime", "alarmname"},
        {"neid", "generationtime"},
        {"siteid", "alarmname", "occurrencetime"},
    )
    header_index = next(
        (
            index for index, row in enumerate(grid)
            if any(required <= {_normalise_header(value) for value in row} for required in required_headers)
        ),
        None,
    )
    if header_index is None:
        raise ValueError(f"{filename}: no recognized alarm header found in CSV.")
    return _rows_from_grid(grid[header_index:])


def _xlsx_member_names(payload: bytes) -> set[str] | None:
    if not zipfile.is_zipfile(BytesIO(payload)):
        return None
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        return {name.casefold() for name in archive.namelist()}


def read_table(payload: bytes, filename: str) -> list[dict[str, Any]]:
    """Read Excel, CSV, or NMS ZIP exports using content signatures."""
    if payload.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        return _read_xls(payload)

    member_names = _xlsx_member_names(payload)
    if member_names is not None:
        if "xl/workbook.xml" in member_names:
            return _read_xlsx(payload)

        with zipfile.ZipFile(BytesIO(payload)) as archive:
            candidates = [
                (name, archive.read(name))
                for name in archive.namelist()
                if not name.endswith("/")
            ]
        parsed = []
        for member_name, member_payload in candidates:
            try:
                rows = read_table(member_payload, member_name)
            except (ValueError, OSError, xlrd.XLRDError, zipfile.BadZipFile):
                continue
            parsed.append((detect_alarm_vendor(rows) is not None, rows))
        if parsed:
            return max(parsed, key=lambda item: item[0])[1]
        raise ValueError(f"{filename}: no alarm workbook or table found in archive.")

    return _read_csv(payload, filename)


def build_site_index(master_rows: list[dict[str, Any]]) -> dict[str, dict[str, str]]:
    sites = {}
    for row in master_rows:
        site_id = _text(_get(row, "Site ID", "Site ID OPSEL")).upper()
        if not site_id:
            continue
        opsel_id = _text(_get(row, "Site ID OPSEL")).upper()
        enodeb_id = _text(_get(row, "enodeb ID/enodeb ID OPSEL", "enodeb ID", "eNodeB ID"))
        info = {
            "site_id": site_id,
            "site_name": _text(_get(row, "Site Name")),
            "vsat_provider": _text(_get(row, "Satellite/MW Brand", "VSAT Provider", "Vsat")),
            "spotbeam": _text(_get(row, "Spotbeam")),
            "status_bapl": _text(_get(row, "Status")),
            "region": _text(_get(row, "Provinsi", "Provinsi Lama")),
            "paket": _text(_get(row, "Paket Penyedia MSS", "Paket")),
            "area": _text(_get(row, "Paket Area Opsel", "Cluster Provinsi")),
            "vendor_transmission": _text(_get(row, "Vendor Transmisi")),
            "enodeb_id": enodeb_id,
        }
        sites[site_id] = info
        if opsel_id and opsel_id not in sites:
            sites[opsel_id] = info
        if enodeb_id and enodeb_id not in sites:
            sites[enodeb_id] = info
    if not sites:
        raise ValueError("Site ID column was not found in Site Management.")
    return sites


def _build_site_id_pattern(site_ids: set[str]) -> re.Pattern[str] | None:
    if not site_ids:
        return None
    alternatives = "|".join(
        re.escape(site_id)
        for site_id in sorted(site_ids, key=len, reverse=True)
    )
    return re.compile(rf"(?<![A-Z0-9])({alternatives})(?![A-Z0-9])")


def _find_site_id(source: Any, pattern: re.Pattern[str] | None, site_ids: set[str] | None = None) -> str:
    text = _text(source).upper()
    if not text:
        return ""
    if site_ids:
        for token in re.findall(r"[A-Z0-9]+", text):
            if token in site_ids:
                return token
    if pattern is not None:
        match = pattern.search(text)
        return match.group(1) if match else ""
    return ""


def _parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = _text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d %H:%M:%S", "%m-%d-%Y %H:%M:%S", "%m/%d/%Y %H:%M:%S"):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).replace(tzinfo=None)
    except ValueError:
        return None


def normalize_alarms(
    alarm_rows: list[dict[str, Any]],
    site_index: dict[str, dict[str, str]],
    vendor: str,
) -> list[dict[str, Any]]:
    """Map ZTE, Fiberhome, Huawei, or already-normalised alarm rows to one schema."""
    site_ids = set(site_index)
    site_id_pattern = _build_site_id_pattern(site_ids)
    result = []
    for row in alarm_rows:
        source = _get(
            row,
            "ME",
            "Fault Source",
            "Alarm Source",
            "Site ID",
            "NE Name",
            "Site",
            "eNodeB ID",
            "ME ID",
            "NE ID",
        )
        if not _text(source):
            source = " ".join(
                _text(_get(row, name))
                for name in ("Additional Information", "Location Information", "MO Name")
            )
        site_id = _text(_get(row, "Site ID")).upper()
        if site_id not in site_ids:
            # Check eNodeB ID directly
            enb = _text(_get(row, "eNodeB ID", "enodeb ID", "ME ID", "NE ID"))
            if enb and enb in site_index:
                site_id = enb
            else:
                site_id = _find_site_id(source, site_id_pattern, site_ids)
        if not site_id:
            continue

        site = site_index[site_id]
        primary_site_id = site.get("site_id") or site_id

        alarm_name = _text(_get(row, "Alarm Code Name", "Alarm Name", "Alarm", "Name"))
        occurrence = _parse_datetime(_get(
            row,
            "Occurrence Time",
            "Generation Time",
            "Start Time",
            "Occurred On (NT)",
            "Occurred On",
            "First Occurred (NT)",
            "Last Occurred (NT)",
        ))
        if not alarm_name or occurrence is None:
            continue
        clear_time = _parse_datetime(_get(
            row,
            "Clear Time",
            "Clear Time (Local)",
            "Cleared On (NT)",
            "Cleared On",
        ))
        clear_status = _text(_get(row, "Clear Status", "Clearance Status")).casefold()
        is_cleared = clear_time is not None or (
            "clear" in clear_status and "unclear" not in clear_status
        )

        duration_str = format_duration(occurrence, clear_time)
        duration_sec = int((clear_time - occurrence).total_seconds()) if clear_time and clear_time >= occurrence else None

        result.append({
            "Vendor": vendor,
            "Site ID": primary_site_id,
            "Site Name": site["site_name"] or _text(_get(row, "Site Name(Office)", "Site Name")),
            "VSAT Provider": site["vsat_provider"] or _text(_get(row, "VSAT Provider", "Vsat")),
            "Spotbeam": site["spotbeam"],
            "Status BAPL": site["status_bapl"] or _text(_get(row, "Status BAPL")),
            "Region": site["region"],
            "Paket": site.get("paket", ""),
            "Area": site.get("area", ""),
            "Vendor Transmisi": site.get("vendor_transmission", ""),
            "Alarm Name": alarm_name,
            "Alarm Type": _text(_get(row, "Alarm Type", "Alarm Category")),
            "Severity": _text(_get(row, "Alarm Severity", "Level", "Severity")),
            "Occurrence Time": occurrence,
            "Clear Time": clear_time,
            "Alarm State": "Cleared" if is_cleared else "Active",
            "Duration": duration_str,
            "Duration Seconds": duration_sec,
            "NE ID": _text(_get(row, "ME ID", "NE ID", "eNodeB ID")),
            "Alarm ID": _text(_get(row, "Alarm ID", "Alarm No.", "Alarm Code")),
            "_raw": dict(row),
        })
    return result


def summarize_providers(
    alarms: list[dict[str, Any]],
    selected_alarm_names: set[str],
    minimum_sites: int = 5,
    group_by_spotbeam: bool = False,
    now: datetime | None = None,
    range_start: datetime | None = None,
    range_end: datetime | None = None,
) -> list[dict[str, Any]]:
    """Return current counts and historical simultaneous-site incidents by VSAT."""
    now = now or datetime.now()
    if range_start is not None and range_end is not None and range_start >= range_end:
        raise ValueError("Start time must be earlier than end time.")

    groups: dict[tuple[str, str], dict[str, Any]] = {}
    for alarm in alarms:
        if alarm["Alarm Name"] not in selected_alarm_names:
            continue
        provider = alarm["VSAT Provider"] or "Unknown"
        spotbeam = alarm["Spotbeam"] or "Unknown"
        key = (provider, spotbeam if group_by_spotbeam else "")
        group = groups.setdefault(key, {"events": {}, "current_sites": set(), "alarms": []})
        group["alarms"].append(alarm)
        site_id = alarm["Site ID"]
        if alarm["Alarm State"] == "Active":
            group["current_sites"].add(site_id)

        occurrence = alarm["Occurrence Time"]
        start = occurrence
        end = alarm["Clear Time"] or now
        if range_start is not None:
            if end <= range_start:
                continue
            start = max(start, range_start)
        if range_end is not None:
            if start >= range_end:
                continue
            end = min(end, range_end)
        if end < start:
            end = start
        if end == start:
            continue
        group["events"].setdefault(start, []).append((1, site_id, occurrence))
        group["events"].setdefault(end, []).append((-1, site_id, occurrence))

    summaries = []

    def incident_sites_at(
        active_starts: dict[str, dict[datetime, int]],
        timestamp: datetime,
    ) -> set[str]:
        window_start = timestamp - timedelta(minutes=45)
        window_end = timestamp + timedelta(minutes=15)
        return {
            site_id
            for site_id, starts in active_starts.items()
            if any(
                count > 0 and window_start <= occurrence <= window_end
                for occurrence, count in starts.items()
            )
        }

    def build_summary(
        provider: str,
        spotbeam: str,
        group: dict[str, Any],
        peak_count: int,
        peak_time: datetime | None,
        peak_sites: set[str],
    ) -> dict[str, Any]:
        return {
            "Range Start": range_start,
            "Range End": range_end,
            "VSAT Provider": provider,
            "Spotbeam": spotbeam,
            "Current Down Sites": len(group["current_sites"]),
            "Peak Simultaneous Sites": peak_count,
            "Peak Time": peak_time,
            "Massive Down": peak_count >= minimum_sites,
            "Sites at Peak": sorted(peak_sites),
        }

    for (provider, spotbeam), group in groups.items():
        active_starts: dict[str, dict[datetime, int]] = {}
        best_count = 0
        best_sites: set[str] = set()
        best_time = None
        open_incident: dict[str, Any] | None = None
        incidents: list[dict[str, Any]] = []

        for timestamp in sorted(group["events"]):
            changes: dict[tuple[str, datetime], int] = {}
            for delta, site_id, occurrence in group["events"][timestamp]:
                key = (site_id, occurrence)
                changes[key] = changes.get(key, 0) + delta
            for (site_id, occurrence), delta in changes.items():
                starts = active_starts.setdefault(site_id, {})
                starts[occurrence] = starts.get(occurrence, 0) + delta
                if starts[occurrence] <= 0:
                    starts.pop(occurrence, None)
                if not starts:
                    active_starts.pop(site_id, None)

            incident_sites = incident_sites_at(active_starts, timestamp)
            incident_count = len(incident_sites)
            if incident_count > best_count:
                best_count = incident_count
                best_sites = set(incident_sites)
                best_time = timestamp

            if incident_count >= minimum_sites:
                if open_incident is None:
                    open_incident = {
                        "peak_count": incident_count,
                        "peak_sites": set(incident_sites),
                        "peak_time": timestamp,
                    }
                elif incident_count > open_incident["peak_count"]:
                    open_incident = {
                        "peak_count": incident_count,
                        "peak_sites": set(incident_sites),
                        "peak_time": timestamp,
                    }
            elif open_incident is not None:
                incidents.append(open_incident)
                open_incident = None

        if open_incident is not None:
            incidents.append(open_incident)

        if incidents:
            for incident in incidents:
                summaries.append(build_summary(
                    provider,
                    spotbeam,
                    group,
                    incident["peak_count"],
                    incident["peak_time"],
                    incident["peak_sites"],
                ))
        else:
            summaries.append(build_summary(provider, spotbeam, group, best_count, best_time, best_sites))

    return sorted(summaries, key=lambda item: (
        -item["Peak Simultaneous Sites"],
        item["VSAT Provider"],
        item["Spotbeam"],
        item["Peak Time"] or datetime.min,
    ))
