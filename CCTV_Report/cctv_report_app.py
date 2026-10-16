from __future__ import annotations

import io
import posixpath
import re
import shutil
import sys
import tempfile
import time
import xml.etree.ElementTree as ET
import zipfile
from bisect import bisect_right
from collections import Counter
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageOps, ImageStat
from PyQt6.QtCore import QEvent, QPointF, QRectF, Qt, QThread, QTimer, pyqtSignal
from PyQt6.QtGui import QColor, QFont, QIcon, QImage, QPainter, QPen, QPixmap
from PyQt6.QtWidgets import (
    QApplication,
    QCheckBox,
    QComboBox,
    QDialog,
    QFileDialog,
    QFrame,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QMainWindow,
    QMessageBox,
    QMenu,
    QPushButton,
    QProgressBar,
    QTableWidget,
    QTableWidgetItem,
    QTabWidget,
    QVBoxLayout,
    QWidget,
)

from noc_brand import (
    DARK_THEME,
    NocLogoMark,
    ThemeSelector,
    palette_for,
    set_application_theme,
    system_theme_name,
)


APP_NAME = "Extract CCTV ScreenShot v.1.0.0"
REPORT_TEMPLATE_FILES = {
    "Paket 10": "Template_Report CCTV Paket 10.xlsx",
    "Paket 11": "Template_Report CCTV Paket 11.xlsx",
}
IMAGE_IMPORT_TEMPLATE_FILES = {
    "Paket 10": "Template_Import_Gambar_CCTV Paket 10.xlsx",
    "Paket 11": "Template_Import_Gambar_CCTV Paket 11.xlsx",
}
EXPORT_FORMAT_PICTURE_IN_CELL = "picture_in_cell"
EXPORT_FORMAT_FLOATING = "floating"
ExportFormat = Literal["picture_in_cell", "floating"]
CAMERA_STATUSES = ("Online", "Offline", "Not Detected", "Online With Note")
CLEAN_STATUSES = ("Bersih", "Kurang Bersih", "Tidak Bersih", "Not Detected")
BOX_LABELS = {
    "panel": ("Sensor Monitoring", QColor("#2f6fed")),
    "cam1": ("Camera 1", QColor("#1f9d55")),
    "cam2": ("Camera 2", QColor("#e08a00")),
    "idbox": ("Site ID", QColor("#8e44ad")),
}
IMAGE_SIZES = {
    "Kecil": (260, 360, 72),
    "Sedang": (360, 520, 78),
    "Besar": (480, 760, 85),
}
PROFILE_KEYS = ("ZTE", "FH", "P11")
PROFILE_LABELS = {
    "ZTE": "ZTE",
    "FH": "FH / Fiberhome",
    "P11": "Paket 11 / HoloSens lama",
}
DEFAULT_CALIBRATIONS = {
    "ZTE": {
        "panel": (0.0, 0.0, 0.215, 0.99),
        "cam1": (0.23, 0.075, 0.375, 0.405),
        "cam2": (0.61, 0.075, 0.38, 0.405),
        "idbox": (0.0, 0.0, 0.02, 0.02),
    },
    "FH": {
        "panel": (0.0, 0.09, 0.23, 0.82),
        "cam1": (0.235, 0.09, 0.38, 0.43),
        "cam2": (0.618, 0.09, 0.38, 0.43),
        "idbox": (0.0, 0.0, 0.02, 0.02),
    },
    "P11": {
        "panel": (0.0, 0.099, 0.167, 0.56),
        "cam1": (0.1737, 0.1025, 0.413, 0.874),
        "cam2": (0.588, 0.1025, 0.412, 0.874),
        "idbox": (0.011, 0.205, 0.088, 0.07),
    },
}
OOXML_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
}


def app_asset_path(name: str) -> Path:
    base_path = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base_path / "assets" / name


def example_template_path(name: str) -> Path:
    base_path = Path(
        getattr(sys, "_MEIPASS", Path(__file__).resolve().parent.parent)
    )
    return base_path / "contoh_template" / name


@dataclass
class ReportItem:
    path: Path
    source_type: str = "FH"
    site_id: str = ""
    site_name: str = ""
    site_name_manual: bool = False
    cam1: str = ""
    cam2: str = ""
    clean: str = ""
    nvr: str = ""
    note: str = ""
    crops: dict[str, bytes] = field(default_factory=dict)
    processed: bool = False
    nvr_manual: bool = False
    manual_fields: set[str] = field(default_factory=set)


@dataclass
class ExtractedCapture:
    path: Path
    site_id: str
    source_type: str
    crop_key: str | None = None
    group_key: str | None = None


def natural_key(path: Path) -> list[Any]:
    return [int(part) if part.isdigit() else part.casefold()
            for part in re.split(r"(\d+)", path.name)]


def site_id_from_filename(path: Path) -> str:
    match = re.search(r"([A-Z]{3}\d{3,5})", path.stem.upper())
    return match.group(1) if match else ""


def derive_nvr(cam1: str, cam2: str) -> str:
    if not cam1 and not cam2:
        return ""
    if cam1 == cam2 and cam1 in ("Offline", "Not Detected"):
        return cam1
    return "Online"


def image_status(image: Image.Image) -> tuple[str, str]:
    sample = image.convert("L").resize((48, 30))
    stats = ImageStat.Stat(sample)
    mean = stats.mean[0]
    deviation = stats.stddev[0]
    if mean > 165 and deviation < 28:
        return "Offline", ""
    if mean < 28 and deviation < 12:
        return "Online With Note", "Video hitam"
    return "Online", ""


def crop_image(image: Image.Image, normalized: tuple[float, float, float, float],
               max_width: int, quality: int) -> bytes:
    x, y, width, height = normalized
    left = round(x * image.width)
    top = round(y * image.height)
    right = max(left + 1, round((x + width) * image.width))
    bottom = max(top + 1, round((y + height) * image.height))
    cropped = image.crop((left, top, right, bottom)).convert("RGB")
    if cropped.width > max_width:
        resized_height = max(1, round(cropped.height * max_width / cropped.width))
        cropped = cropped.resize((max_width, resized_height), Image.Resampling.LANCZOS)
    output = io.BytesIO()
    cropped.save(output, format="JPEG", quality=quality, optimize=True)
    return output.getvalue()


class CalibrationCanvas(QWidget):
    calibration_changed = pyqtSignal()

    def __init__(self, profile: str = "FH",
                 parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setMinimumSize(760, 340)
        self.setMouseTracking(True)
        self.image: QPixmap | None = None
        self.profile = profile
        self.boxes = dict(DEFAULT_CALIBRATIONS[profile])
        self._active: str | None = None
        self._resize = False
        self._last = QPointF()
        self._canvas_background = QColor("#172B48")
        self._canvas_text = QColor("#A8B9CF")

    def set_theme_colors(self, background: str, text: str) -> None:
        self._canvas_background = QColor(background)
        self._canvas_text = QColor(text)
        self.update()

    def set_image(self, path: Path | None) -> None:
        pixmap = QPixmap(str(path)) if path is not None else QPixmap()
        self.image = pixmap if not pixmap.isNull() else None
        self.update()

    def reset_boxes(self) -> None:
        self.boxes = dict(DEFAULT_CALIBRATIONS[self.profile])
        self.calibration_changed.emit()
        self.update()

    def normalized_boxes(self) -> dict[str, tuple[float, float, float, float]]:
        return dict(self.boxes)

    def _image_rect(self) -> QRectF:
        bounds = QRectF(self.rect()).adjusted(8, 8, -8, -8)
        if not self.image or self.image.width() <= 0 or self.image.height() <= 0:
            return bounds
        scale = min(bounds.width() / self.image.width(),
                    bounds.height() / self.image.height())
        width = self.image.width() * scale
        height = self.image.height() * scale
        return QRectF(bounds.center().x() - width / 2,
                      bounds.center().y() - height / 2, width, height)

    def paintEvent(self, event: Any) -> None:
        del event
        painter = QPainter(self)
        painter.fillRect(self.rect(), self._canvas_background)
        image_rect = self._image_rect()
        if self.image:
            painter.drawPixmap(image_rect, self.image,
                               QRectF(self.image.rect()))
        else:
            painter.setPen(self._canvas_text)
            painter.drawText(self.rect(), Qt.AlignmentFlag.AlignCenter,
                             "Tambahkan screenshot untuk mengatur area crop")
            return
        for key, (label, color) in BOX_LABELS.items():
            x, y, width, height = self.boxes[key]
            rect = QRectF(image_rect.x() + x * image_rect.width(),
                          image_rect.y() + y * image_rect.height(),
                          width * image_rect.width(),
                          height * image_rect.height())
            painter.setPen(QPen(color, 2))
            fill = QColor(color)
            fill.setAlpha(38)
            painter.fillRect(rect, fill)
            painter.drawRect(rect)
            painter.fillRect(QRectF(rect.x(), rect.y(), min(190, rect.width()), 21),
                             color)
            painter.setPen(Qt.GlobalColor.white)
            painter.setFont(QFont("Segoe UI", 9, QFont.Weight.DemiBold))
            painter.drawText(QPointF(rect.x() + 5, rect.y() + 15), label)
            painter.fillRect(QRectF(rect.right() - 8, rect.bottom() - 8, 12, 12),
                             color)

    def mousePressEvent(self, event: Any) -> None:
        image_rect = self._image_rect()
        point = event.position()
        for key in reversed(tuple(BOX_LABELS)):
            x, y, width, height = self.boxes[key]
            rect = QRectF(image_rect.x() + x * image_rect.width(),
                          image_rect.y() + y * image_rect.height(),
                          width * image_rect.width(),
                          height * image_rect.height())
            if QRectF(rect.right() - 12, rect.bottom() - 12, 18, 18).contains(point):
                self._active, self._resize = key, True
                break
            if rect.contains(point):
                self._active, self._resize = key, False
                break
        self._last = point

    def mouseMoveEvent(self, event: Any) -> None:
        if not self._active:
            return
        image_rect = self._image_rect()
        if image_rect.width() <= 0 or image_rect.height() <= 0:
            return
        dx = (event.position().x() - self._last.x()) / image_rect.width()
        dy = (event.position().y() - self._last.y()) / image_rect.height()
        x, y, width, height = self.boxes[self._active]
        if self._resize:
            width = min(1 - x, max(0.02, width + dx))
            height = min(1 - y, max(0.02, height + dy))
        else:
            x = min(1 - width, max(0, x + dx))
            y = min(1 - height, max(0, y + dy))
        self.boxes[self._active] = (x, y, width, height)
        self._last = event.position()
        self.calibration_changed.emit()
        self.update()

    def mouseReleaseEvent(self, event: Any) -> None:
        del event
        self._active = None


class ProcessIndicator(QWidget):
    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setFixedSize(24, 24)
        self._state = "idle"
        self._angle = 0
        self._progress = 0.0
        self._timer = QTimer(self)
        self._timer.setInterval(30)
        self._timer.timeout.connect(self._advance_animation)

    def set_running(self) -> None:
        self._state = "running"
        self._progress = 0.0
        self._timer.start()
        self.update()

    def set_success(self) -> None:
        self._state = "success"
        self._progress = 0.0
        self._timer.start()
        self.update()

    def set_error(self) -> None:
        self._timer.stop()
        self._state = "error"
        self._progress = 1.0
        self.update()

    def set_idle(self) -> None:
        self._timer.stop()
        self._state = "idle"
        self._progress = 0.0
        self.hide()
        self.update()

    def _advance_animation(self) -> None:
        if self._state == "running":
            self._angle = (self._angle + 12) % 360
        elif self._state == "success":
            self._progress = min(1.0, self._progress + 0.08)
            if self._progress == 1.0:
                self._timer.stop()
        self.update()

    def paintEvent(self, event: Any) -> None:
        del event
        if self._state == "idle":
            return
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        bounds = QRectF(2, 2, 20, 20)
        if self._state == "running":
            painter.setPen(QPen(QColor("#34434c"), 2.5))
            painter.drawEllipse(bounds)
            painter.setPen(QPen(
                QColor("#42d5be"), 2.5, Qt.PenStyle.SolidLine,
                Qt.PenCapStyle.RoundCap,
            ))
            painter.drawArc(bounds, self._angle * 16, 105 * 16)
        elif self._state == "success":
            radius = 3 + 7 * self._progress
            center = bounds.center()
            painter.setPen(Qt.PenStyle.NoPen)
            painter.setBrush(QColor("#35b779"))
            painter.drawEllipse(center, radius, radius)
            if self._progress > 0.55:
                check_progress = min(1.0, (self._progress - 0.55) / 0.45)
                painter.setPen(QPen(
                    QColor("#10251c"), 2.2, Qt.PenStyle.SolidLine,
                    Qt.PenCapStyle.RoundCap, Qt.PenJoinStyle.RoundJoin,
                ))
                first = QPointF(7, 12)
                corner = QPointF(10, 15)
                end = QPointF(17, 8)
                if check_progress < 0.5:
                    middle = first + (corner - first) * (check_progress * 2)
                    painter.drawLine(first, middle)
                else:
                    painter.drawLine(first, corner)
                    middle = corner + (end - corner) * (
                        (check_progress - 0.5) * 2
                    )
                    painter.drawLine(corner, middle)
        else:
            painter.setPen(Qt.PenStyle.NoPen)
            painter.setBrush(QColor("#e25d5d"))
            painter.drawEllipse(bounds)
            painter.setPen(QPen(
                QColor("#ffffff"), 2.4, Qt.PenStyle.SolidLine,
                Qt.PenCapStyle.RoundCap,
            ))
            painter.drawLine(QPointF(12, 7), QPointF(12, 13))
            painter.drawPoint(QPointF(12, 17))


class ProcessingWorker(QThread):
    progress = pyqtSignal(int, int, str)
    item_finished = pyqtSignal(int, object, str)

    def __init__(self, items: list[ReportItem],
                 calibrations: dict[
                     str, dict[str, tuple[float, float, float, float]]
                 ],
                 panel_width: int, camera_width: int, quality: int) -> None:
        super().__init__()
        self.items = items
        self.calibrations = calibrations
        self.panel_width = panel_width
        self.camera_width = camera_width
        self.quality = quality

    def run(self) -> None:
        for index, item in enumerate(self.items):
            try:
                calibration = self.calibrations.get(
                    item.source_type, DEFAULT_CALIBRATIONS["FH"]
                )
                with Image.open(item.path) as source:
                    image = ImageOps.exif_transpose(source).convert("RGB")
                    crops = {
                        key: crop_image(image, calibration[key],
                                        self.panel_width if key == "panel"
                                        else self.camera_width,
                                        self.quality)
                        for key in ("panel", "cam1", "cam2")
                    }
                    statuses = {}
                    notes = []
                    for key in ("cam1", "cam2"):
                        with Image.open(io.BytesIO(crops[key])) as crop:
                            status, note = image_status(crop)
                        statuses[key] = status
                        if note:
                            notes.append(f"Camera {1 if key == 'cam1' else 2}: {note}")
                    result = {"crops": crops, "statuses": statuses,
                              "note": "; ".join(notes)}
                self.item_finished.emit(index, result, "")
            except (OSError, ValueError) as error:
                self.item_finished.emit(index, None, str(error))
            self.progress.emit(index + 1, len(self.items), item.path.name)


def find_site_columns(workbook: Any) -> tuple[Any, int, int, int]:
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            headers = {
                str(cell.value).strip().casefold(): cell.column
                for cell in row if cell.value is not None
            }
            if "site id" in headers and "site name" in headers:
                return worksheet, row[0].row, headers["site id"], headers["site name"]
    raise ValueError('Kolom "Site ID" dan "Site Name" tidak ditemukan.')


def normalize_status(value: Any, camera: bool = False) -> str:
    text = str(value or "").strip()
    upper = text.upper()
    if camera and upper == "OK":
        return "Online"
    if camera and upper == "NOK":
        return "Offline"
    return {
        "ONLINE": "Online",
        "OFFLINE": "Offline",
        "NOT DETECTED": "Not Detected",
        "ONLINE WITH NOTE": "Online With Note",
    }.get(upper, "")


def find_camera_status_column(headers: dict[str, int], camera: str,
                              location: str) -> int | None:
    candidates = [(header, column) for header, column in headers.items()
                  if header.startswith(camera)]
    preferred = next(
        (column for header, column in candidates
         if any(term in header for term in ("blt", "ptz", location))),
        None,
    )
    if preferred:
        return preferred
    return next(
        (column for header, column in candidates
         if header == camera or "(" in header),
        candidates[0][1] if candidates else None,
    )


def load_site_data(path: Path) -> tuple[dict[str, str], dict[str, int],
                                       dict[str, dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet, header_row, id_column, name_column = find_site_columns(workbook)
        headers = {
            re.sub(r"\s+", " ", str(cell.value or "")).strip().casefold(): cell.column
            for cell in worksheet[header_row]
        }
        cam1_column = find_camera_status_column(headers, "camera 1", "kabinet")
        cam2_column = find_camera_status_column(headers, "camera 2", "pintu pagar")
        columns = {
            "nvr": headers.get("nvr"),
            "cam1": cam1_column,
            "cam2": cam2_column,
            "clean": headers.get("kebersihan"),
            "note": headers.get("keterangan"),
            "provider": headers.get("konsorsium pembangunan"),
        }
        site_map: dict[str, str] = {}
        site_order: dict[str, int] = {}
        site_fields: dict[str, dict[str, Any]] = {}
        for row in worksheet.iter_rows(min_row=header_row + 1):
            site_id = str(row[id_column - 1].value or "").strip().upper()
            if not site_id:
                continue
            site_name = str(row[name_column - 1].value or "").strip()
            site_map[site_id] = site_name
            site_order.setdefault(site_id, len(site_order))
            fields: dict[str, Any] = {}
            for key, column in columns.items():
                if column is None or column > len(row):
                    continue
                value = row[column - 1].value
                if key in ("cam1", "cam2", "nvr"):
                    value = normalize_status(value, camera=key in ("cam1", "cam2"))
                elif value is not None:
                    value = str(value).strip()
                if value:
                    fields[key] = value
            site_fields[site_id] = fields
        return site_map, site_order, site_fields
    finally:
        workbook.close()


def find_duplicate_site_rows(path: Path) -> list[tuple[str, str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet, header_row, id_column, name_column = find_site_columns(workbook)
        sites: dict[str, tuple[str, int]] = {}
        for row in worksheet.iter_rows(min_row=header_row + 1):
            site_id = str(row[id_column - 1].value or "").strip().upper()
            if not site_id:
                continue
            site_name = str(row[name_column - 1].value or "").strip()
            previous_name, count = sites.get(site_id, (site_name, 0))
            sites[site_id] = (previous_name or site_name, count + 1)
        return [
            (site_id, site_name, count)
            for site_id, (site_name, count) in sites.items()
            if count > 1
        ]
    finally:
        workbook.close()


def resolve_ooxml_target(source_part: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def capture_profile(sheet_name: str, providers: list[str]) -> str:
    name = re.sub(r"[^A-Z]", "", sheet_name.upper())
    if name.startswith("ZTE"):
        return "ZTE"
    if name.startswith(("FH", "FIBERHOME")):
        return "FH"
    if name == "CAPTURE" and any(
        "LA-HW-SEI" in provider.upper() for provider in providers
    ):
        return "P11"
    zte_count = sum("ZTE" in provider.upper() for provider in providers)
    fiberhome_count = sum(
        "FIBERHOME" in provider.upper() or "FH-" in provider.upper()
        for provider in providers
    )
    return "ZTE" if zte_count > fiberhome_count else "FH"


def read_ooxml_cell_text(cell: ET.Element, shared_strings: list[str],
                         namespaces: dict[str, str]) -> str:
    if cell.attrib.get("t") == "inlineStr":
        return "".join(
            text.text or "" for text in cell.findall(".//main:t", namespaces)
        ).strip()
    value = cell.find("main:v", namespaces)
    if value is None:
        return ""
    if cell.attrib.get("t") == "s":
        index = int(value.text or "0")
        return shared_strings[index].strip() if index < len(shared_strings) else ""
    return (value.text or "").strip()


def find_rich_image_media(
    archive: zipfile.ZipFile, names: set[str]
) -> dict[int, str]:
    namespaces = {
        **OOXML_NS,
        "rd": "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata",
        "rvrel": "http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel",
    }
    workbook_part = "xl/workbook.xml"
    workbook_rels_part = "xl/_rels/workbook.xml.rels"
    if workbook_part not in names or workbook_rels_part not in names:
        return {}
    workbook_rels = ET.fromstring(archive.read(workbook_rels_part))
    targets = {
        rel.attrib["Type"].rsplit("/", 1)[-1]: resolve_ooxml_target(
            workbook_part, rel.attrib["Target"]
        )
        for rel in workbook_rels.findall("rel:Relationship", namespaces)
    }
    metadata_part = targets.get("sheetMetadata")
    rich_value_rel_part = targets.get("richValueRel")
    if not metadata_part or not rich_value_rel_part:
        return {}
    rich_values_part = posixpath.join(
        posixpath.dirname(rich_value_rel_part), "rdrichvalue.xml"
    )
    relationships_part = posixpath.join(
        posixpath.dirname(rich_value_rel_part), "_rels",
        posixpath.basename(rich_value_rel_part) + ".rels",
    )
    if not all(part in names for part in (
        metadata_part, rich_value_rel_part, rich_values_part, relationships_part
    )):
        return {}

    metadata_root = ET.fromstring(archive.read(metadata_part))
    future_metadata = metadata_root.find(
        "main:futureMetadata[@name='XLRICHVALUE']", namespaces
    )
    value_metadata = metadata_root.find("main:valueMetadata", namespaces)
    if future_metadata is None or value_metadata is None:
        return {}
    rich_indexes: list[int | None] = []
    for entry in future_metadata.findall("main:bk", namespaces):
        rich_value = entry.find(".//rd:rvb", namespaces)
        rich_indexes.append(
            int(rich_value.attrib["i"])
            if rich_value is not None and "i" in rich_value.attrib else None
        )

    value_to_rich: dict[int, int] = {}
    for value_index, entry in enumerate(value_metadata.findall("main:bk", namespaces)):
        reference = entry.find("main:rc", namespaces)
        if reference is None or reference.attrib.get("t") != "1":
            continue
        future_index = int(reference.attrib.get("v", "-1"))
        if 0 <= future_index < len(rich_indexes):
            rich_index = rich_indexes[future_index]
            if rich_index is not None:
                value_to_rich[value_index + 1] = rich_index

    rich_values_root = ET.fromstring(archive.read(rich_values_part))
    rich_to_image: dict[int, int] = {}
    for rich_index, value in enumerate(rich_values_root.findall("rd:rv", namespaces)):
        local_image_id = value.find("rd:v", namespaces)
        if local_image_id is not None and local_image_id.text is not None:
            rich_to_image[rich_index] = int(local_image_id.text)

    rich_rel_root = ET.fromstring(archive.read(rich_value_rel_part))
    relationship_ids = [
        rel.attrib[f"{{{namespaces['r']}}}id"]
        for rel in rich_rel_root.findall("rvrel:rel", namespaces)
    ]
    picture_rels = ET.fromstring(archive.read(relationships_part))
    image_targets = {
        rel.attrib["Id"]: resolve_ooxml_target(
            rich_value_rel_part, rel.attrib["Target"]
        )
        for rel in picture_rels.findall("rel:Relationship", namespaces)
        if rel.attrib.get("Type", "").endswith("/image")
    }
    return {
        value_index: image_targets[relationship_ids[image_index]]
        for value_index, rich_index in value_to_rich.items()
        if (image_index := rich_to_image.get(rich_index)) is not None
        and image_index < len(relationship_ids)
        and relationship_ids[image_index] in image_targets
    }


def extract_embedded_captures(
    workbook_path: Path,
    output_dir: Path,
    site_providers: dict[str, str] | None = None,
    progress_callback: Any | None = None,
) -> list[ExtractedCapture]:
    namespaces = OOXML_NS
    with zipfile.ZipFile(workbook_path) as archive:
        names = set(archive.namelist())
        workbook_part = "xl/workbook.xml"
        workbook_root = ET.fromstring(archive.read(workbook_part))
        workbook_rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        workbook_targets = {
            relation.attrib["Id"]: relation.attrib["Target"]
            for relation in workbook_rels.findall("rel:Relationship", namespaces)
        }
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in names:
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            shared_strings = [
                "".join(text.text or "" for text in item.findall(".//main:t", namespaces))
                for item in shared_root.findall("main:si", namespaces)
            ]
        rich_image_media = find_rich_image_media(archive, names)
        output_dir.mkdir(parents=True, exist_ok=True)
        capture_tasks: list[
            tuple[str, int, str, str, str, str | None, str | None]
        ] = []
        for sheet in workbook_root.findall(".//main:sheets/main:sheet", namespaces):
            sheet_name = sheet.attrib["name"].strip()
            relationship_id = sheet.attrib[f"{{{namespaces['r']}}}id"]
            sheet_part = resolve_ooxml_target(
                workbook_part, workbook_targets[relationship_id]
            )
            if sheet_part not in names:
                continue
            sheet_root = ET.fromstring(archive.read(sheet_part))
            site_ids_by_row: dict[int, str] = {}
            site_id_column = 1
            image_columns: dict[int, str] = {}
            rich_cells: list[tuple[int, str, str]] = []
            sheet_data = sheet_root.find("main:sheetData", namespaces)
            if sheet_data is not None:
                for row in sheet_data.findall("main:row", namespaces):
                    row_number = int(row.attrib["r"])
                    row_cells = row.findall("main:c", namespaces)
                    if sheet_name.casefold() == "site list":
                        headers = {
                            re.sub(r"\s+", " ", read_ooxml_cell_text(
                                cell, shared_strings, namespaces
                            )).casefold(): column_index_from_string(re.match(
                                r"[A-Z]+", cell.attrib["r"]
                            ).group())
                            for cell in row_cells
                            if cell.attrib.get("r")
                            and re.match(r"[A-Z]+", cell.attrib["r"])
                        }
                        if "site id" in headers and "site name" in headers:
                            site_id_column = headers["site id"]
                            image_columns = {
                                headers[header]: crop_key
                                for header, crop_key in (
                                    ("sensor monitoring", "panel"),
                                    ("camera 1", "cam1"),
                                    ("camera 2", "cam2"),
                                )
                                if header in headers
                            }
                    for cell in row_cells:
                        cell_reference = cell.attrib.get("r", "")
                        cell_column = re.match(r"[A-Z]+", cell_reference)
                        if (cell_column and column_index_from_string(
                                cell_column.group()) == site_id_column):
                            site_id = read_ooxml_cell_text(
                                cell, shared_strings, namespaces
                            ).upper()
                            if site_id:
                                site_ids_by_row[row_number] = site_id
                            break
                    for cell in row_cells:
                        metadata_index = cell.attrib.get("vm")
                        if not metadata_index:
                            continue
                        media_part = rich_image_media.get(
                            int(metadata_index)
                        )
                        if media_part:
                            rich_cells.append((
                                row_number, cell.attrib["r"], media_part
                            ))

            provider_values = [
                site_providers.get(site_id, "")
                for site_id in site_ids_by_row.values()
                if site_providers
            ]
            source_type = capture_profile(
                "Capture" if sheet_name.casefold() == "site list" else sheet_name,
                provider_values,
            )

            def profile_for_site(site_id: str) -> str:
                if sheet_name.casefold() != "site list":
                    return source_type
                provider = (site_providers or {}).get(site_id, "")
                return capture_profile("Capture", [provider] if provider else [])

            if sheet_name.casefold() == "site list":
                for row_number, cell_reference, media_part in rich_cells:
                    column_match = re.match(r"[A-Z]+", cell_reference)
                    if not column_match:
                        continue
                    crop_key = image_columns.get(
                        column_index_from_string(column_match.group())
                    )
                    site_id = site_ids_by_row.get(row_number, "")
                    if crop_key and site_id:
                        capture_tasks.append((
                            sheet_name, row_number, media_part, site_id,
                            crop_key, profile_for_site(site_id),
                            f"{sheet_name}:{row_number}",
                        ))
            else:
                site_id_rows = sorted(site_ids_by_row)
                for row_number, _, media_part in rich_cells:
                    preceding = bisect_right(site_id_rows, row_number) - 1
                    site_id = (
                        site_ids_by_row[site_id_rows[preceding]]
                        if preceding >= 0 else ""
                    )
                    capture_tasks.append((
                        sheet_name, row_number, media_part, site_id, "",
                        source_type, None,
                    ))

            sheet_rel_part = posixpath.join(
                posixpath.dirname(sheet_part), "_rels",
                posixpath.basename(sheet_part) + ".rels",
            )
            if sheet_rel_part not in names:
                continue
            sheet_rels = ET.fromstring(archive.read(sheet_rel_part))
            drawing_relation = next(
                (relation for relation in sheet_rels.findall(
                    "rel:Relationship", namespaces
                ) if relation.attrib.get("Type", "").endswith("/drawing")),
                None,
            )
            if drawing_relation is None:
                continue
            drawing_part = resolve_ooxml_target(
                sheet_part, drawing_relation.attrib["Target"]
            )
            drawing_rels_part = posixpath.join(
                posixpath.dirname(drawing_part), "_rels",
                posixpath.basename(drawing_part) + ".rels",
            )
            if drawing_part not in names or drawing_rels_part not in names:
                continue
            drawing_rels = ET.fromstring(archive.read(drawing_rels_part))
            drawing_targets = {
                relation.attrib["Id"]: resolve_ooxml_target(
                    drawing_part, relation.attrib["Target"]
                )
                for relation in drawing_rels.findall(
                    "rel:Relationship", namespaces
                )
            }

            drawing_root = ET.fromstring(archive.read(drawing_part))
            for anchor in list(drawing_root):
                origin = anchor.find("xdr:from", namespaces)
                picture = anchor.find("xdr:pic", namespaces)
                if origin is None or picture is None:
                    continue
                blip = picture.find(".//a:blip", namespaces)
                if blip is None:
                    continue
                embed_id = blip.attrib.get(f"{{{namespaces['r']}}}embed")
                media_part = drawing_targets.get(embed_id or "")
                if media_part not in names:
                    continue
                origin_row = int(origin.find("xdr:row", namespaces).text or "0")
                origin_col = int(origin.find("xdr:col", namespaces).text or "0")
                if sheet_name.casefold() == "site list":
                    crop_key = image_columns.get(origin_col + 1)
                    site_id = site_ids_by_row.get(origin_row + 1, "")
                    if crop_key and site_id:
                        capture_tasks.append((
                            sheet_name, origin_row + 1, media_part, site_id,
                            crop_key, profile_for_site(site_id),
                            f"{sheet_name}:{origin_row + 1}",
                        ))
                    continue
                site_id_rows = sorted(site_ids_by_row)
                preceding = bisect_right(site_id_rows, origin_row + 1) - 1
                site_id = (
                    site_ids_by_row[site_id_rows[preceding]]
                    if preceding >= 0 else ""
                )
                capture_tasks.append((
                    sheet_name, origin_row, media_part, site_id, "",
                    source_type, None,
                ))

        total_anchors = len(capture_tasks)
        if not total_anchors:
            raise ValueError(
                "Tidak ditemukan gambar floating atau Picture in Cell yang "
                "terhubung ke Site ID pada workbook."
            )
        captures: list[ExtractedCapture] = []
        for index, (
            sheet_name, row_number, media_part, site_id, crop_key,
            source_type, group_key,
        ) in enumerate(capture_tasks, start=1):
            safe_id = re.sub(r"[^A-Z0-9_-]", "_", site_id) or "TANPA_ID"
            safe_sheet = re.sub(r"[^A-Za-z0-9_-]", "_", source_type or "FH")
            suffix = Path(media_part).suffix.lower() or ".png"
            crop_suffix = f"_{crop_key}" if crop_key else ""
            image_path = output_dir / (
                f"{index:04d}_{safe_sheet}_{safe_id}{crop_suffix}{suffix}"
            )
            with archive.open(media_part) as source, image_path.open("wb") as target:
                shutil.copyfileobj(source, target)
            captures.append(ExtractedCapture(
                image_path, site_id, source_type or "FH", crop_key or None,
                group_key,
            ))
            if progress_callback:
                progress_callback(index, total_anchors, site_id)
        return captures


class CaptureImportWorker(QThread):
    progress = pyqtSignal(int, int, str)
    completed = pyqtSignal(object, str)

    def __init__(self, workbook_path: Path, output_dir: Path) -> None:
        super().__init__()
        self.workbook_path = workbook_path
        self.output_dir = output_dir

    def run(self) -> None:
        try:
            site_map, site_order, site_fields = load_site_data(self.workbook_path)
            captures = extract_embedded_captures(
                self.workbook_path, self.output_dir,
                {site_id: fields.get("provider", "")
                 for site_id, fields in site_fields.items()},
                lambda done, total, site_id: self.progress.emit(done, total, site_id),
            )
            self.completed.emit({
                "path": self.workbook_path,
                "site_map": site_map,
                "site_order": site_order,
                "site_fields": site_fields,
                "captures": captures,
            }, "")
        except (OSError, ValueError, KeyError, zipfile.BadZipFile, ET.ParseError) as error:
            self.completed.emit(None, str(error))


class ExportWorker(QThread):
    progress = pyqtSignal(int, int, str)
    completed = pyqtSignal(object, str)

    def __init__(self, window: Any, output_path: Path,
                 records: list[ReportItem],
                 image_placement: ExportFormat = EXPORT_FORMAT_PICTURE_IN_CELL) -> None:
        super().__init__()
        self.window = window
        self.output_path = output_path
        self.records = records
        self.image_placement = image_placement

    def run(self) -> None:
        initialized_com = False
        try:
            if self.image_placement == EXPORT_FORMAT_PICTURE_IN_CELL:
                import pythoncom

                pythoncom.CoInitialize()
                initialized_com = True
            else:
                self.progress.emit(
                    0, 0, "Menyimpan data dan gambar floating..."
                )
            if self.image_placement == EXPORT_FORMAT_PICTURE_IN_CELL:
                self.progress.emit(
                    0, 0, "Menyimpan data dan gambar ke template Excel..."
                )
            image_count = self.window._write_report(
                self.output_path, self.records
            )
            if self.image_placement == EXPORT_FORMAT_PICTURE_IN_CELL:
                counts = self.window._place_images_in_cells(
                    self.output_path,
                    lambda done, total, message: self.progress.emit(
                        done, total, message
                    ),
                )
            else:
                counts = (0, image_count, "")
            self.completed.emit(
                (self.output_path, counts, self.image_placement), ""
            )
        except Exception as error:
            self.completed.emit(None, str(error))
        finally:
            if initialized_com:
                pythoncom.CoUninitialize()


class MainWindow(QMainWindow):
    table_headers = (
        "No", "Site ID", "Site Name", "Sensor Monitoring", "Camera 1", "Camera 2",
        "Status C1", "Status C2", "Kebersihan", "NVR", "Keterangan",
    )

    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle(APP_NAME)
        icon_path = app_asset_path("cctv-report-icon.ico")
        if icon_path.is_file():
            self.setWindowIcon(QIcon(str(icon_path)))
        self.resize(1540, 1000)
        self.setMinimumSize(1000, 700)
        self.template_path: Path | None = None
        self.site_map: dict[str, str] = {}
        self.site_order: dict[str, int] = {}
        self.site_fields: dict[str, dict[str, Any]] = {}
        self.items: list[ReportItem] = []
        self.worker: ProcessingWorker | None = None
        self.import_worker: CaptureImportWorker | None = None
        self.export_worker: ExportWorker | None = None
        self.busy = False
        self._last_export_progress_at = 0.0
        self._last_export_progress_message = ""
        self._loading_table = False
        self._clipboard_tempdir = tempfile.TemporaryDirectory(prefix="cctv-report-")
        self._embedded_tempdir = tempfile.TemporaryDirectory(prefix="cctv-captures-")
        self._clipboard_counter = 0
        self._capture_import_counter = 0
        self._imported_capture_sources: set[Path] = set()
        application = QApplication.instance()
        if application:
            application.installEventFilter(self)
        self.current_theme_name = (
            application.property("nocTheme") if application else None
        ) or system_theme_name(application)
        self._build_ui()
        self._apply_theme(self.current_theme_name)

    def _build_ui(self) -> None:
        central = QWidget()
        root = QVBoxLayout(central)
        root.setContentsMargins(18, 14, 18, 16)
        root.setSpacing(12)

        header = QFrame()
        header.setObjectName("cctvHeader")
        header_layout = QHBoxLayout(header)
        header_layout.setContentsMargins(16, 10, 16, 10)
        header_layout.setSpacing(12)
        header_layout.addWidget(NocLogoMark(38))
        title_stack = QVBoxLayout()
        title_stack.setSpacing(2)
        title = QLabel("CCTV Extractor")
        title.setObjectName("headerTitle")
        title_stack.addWidget(title)
        subtitle = QLabel("Kalibrasi capture, tinjau data kamera, dan ekspor laporan")
        subtitle.setObjectName("brandSubtitle")
        title_stack.addWidget(subtitle)
        header_layout.addLayout(title_stack, 1)
        self.theme_selector = ThemeSelector(self.current_theme_name)
        self.theme_selector.themeChanged.connect(self._apply_theme)
        header_layout.addWidget(self.theme_selector)
        root.addWidget(header)

        self.busy_banner = QWidget()
        self.busy_banner.setObjectName("busyBanner")
        busy_layout = QHBoxLayout(self.busy_banner)
        busy_layout.setContentsMargins(10, 6, 10, 6)
        busy_layout.setSpacing(10)
        self.activity_indicator = ProcessIndicator()
        self.busy_status_label = QLabel()
        self.busy_status_label.setWordWrap(True)
        self.busy_progress = QProgressBar()
        self.busy_progress.setRange(0, 0)
        self.busy_progress.setMaximumWidth(190)
        busy_layout.addWidget(self.activity_indicator)
        busy_layout.addWidget(self.busy_status_label, 1)
        busy_layout.addWidget(self.busy_progress)
        self.export_watchdog = QTimer(self)
        self.export_watchdog.setInterval(1000)
        self.export_watchdog.timeout.connect(self._update_export_watchdog)
        root.addWidget(self.busy_banner)
        self.busy_banner.hide()

        self.main_tabs = QTabWidget()
        root.addWidget(self.main_tabs)
        prepare_page = QWidget()
        prepare_layout = QVBoxLayout(prepare_page)
        prepare_layout.setContentsMargins(8, 8, 8, 8)
        prepare_layout.setSpacing(10)

        sources = QGroupBox("1. Sumber data")
        source_layout = QHBoxLayout(sources)
        self.template_label = QLabel("Template Excel belum dipilih")
        self.template_button = QPushButton("Pilih template .xlsx")
        self.template_button.clicked.connect(self.choose_template)
        self.download_report_template_button = QPushButton(
            "Unduh template laporan"
        )
        self.download_report_template_button.setToolTip(
            "Unduh contoh template laporan untuk Paket 10 atau Paket 11."
        )
        report_template_menu = QMenu(self)
        for package, filename in REPORT_TEMPLATE_FILES.items():
            action = report_template_menu.addAction(package)
            action.triggered.connect(
                lambda _checked=False, name=filename: self.download_example_template(name)
            )
        self.download_report_template_button.setMenu(report_template_menu)
        self.import_workbook_button = QPushButton("Import gambar dari Excel")
        self.import_workbook_button.clicked.connect(self.import_capture_workbook)
        self.download_image_template_button = QPushButton(
            "Unduh template import gambar"
        )
        self.download_image_template_button.setToolTip(
            "Unduh contoh workbook gambar tertanam untuk Paket 10 atau Paket 11."
        )
        image_template_menu = QMenu(self)
        for package, filename in IMAGE_IMPORT_TEMPLATE_FILES.items():
            action = image_template_menu.addAction(package)
            action.triggered.connect(
                lambda _checked=False, name=filename: self.download_example_template(name)
            )
        self.download_image_template_button.setMenu(image_template_menu)
        self.images_label = QLabel("Belum ada screenshot")
        self.images_button = QPushButton("Pilih screenshot")
        self.images_button.clicked.connect(self.choose_images)
        self.paste_button = QPushButton("Paste gambar (Ctrl+V)")
        self.paste_button.clicked.connect(self.paste_clipboard)
        self.reset_data_button = QPushButton("Reset data")
        self.reset_data_button.clicked.connect(self.reset_all_data)
        source_layout.addWidget(self.template_button)
        source_layout.addWidget(self.download_report_template_button)
        source_layout.addWidget(self.template_label, 1)
        source_layout.addWidget(self.import_workbook_button)
        source_layout.addWidget(self.download_image_template_button)
        source_layout.addWidget(self.images_button)
        source_layout.addWidget(self.paste_button)
        source_layout.addWidget(self.images_label, 1)
        source_layout.addWidget(self.reset_data_button)

        calibration_group = QGroupBox("2. Kalibrasi crop")
        calibration_layout = QVBoxLayout(calibration_group)
        instructions = QLabel(
            "Pilih tab ZTE atau FH / Fiberhome untuk mengatur masing-masing "
            "tata letak. Geser kotak crop dan tarik handle di pojok kanan "
            "bawah untuk mengubah ukurannya. Kotak ungu hanya panduan posisi "
            "Site ID; ID tidak dibaca lewat OCR."
        )
        instructions.setWordWrap(True)
        calibration_layout.addWidget(instructions)
        self.calibration_tabs = QTabWidget()
        self.calibration_canvases = {
            profile: CalibrationCanvas(profile) for profile in PROFILE_KEYS
        }
        for profile in PROFILE_KEYS:
            self.calibration_tabs.addTab(
                self.calibration_canvases[profile], PROFILE_LABELS[profile]
            )
        calibration_layout.addWidget(self.calibration_tabs, 1)

        self.size_combo = QComboBox()
        self.size_combo.addItems(IMAGE_SIZES)
        self.size_combo.setCurrentText("Sedang")
        self.reset_crop_button = QPushButton("Reset kotak crop")
        self.reset_crop_button.clicked.connect(self.reset_current_calibration)
        self.process_button = QPushButton("Proses screenshot")
        self.process_button.setObjectName("accentBtn")
        self.process_button.clicked.connect(self.process_images)
        self.reprocess_button = QPushButton("Proses ulang semua")
        self.reprocess_button.clicked.connect(self.reprocess_images)
        self.progress = QProgressBar()
        self.progress.setValue(0)
        self.process_status = QLabel(
            "Site ID diambil dari nama file atau sheet Capture, ZTE, dan FH; "
            "ID dapat dikoreksi manual pada tabel."
        )
        self.process_status.setWordWrap(True)
        status_layout = QHBoxLayout()
        status_layout.addWidget(self.process_status, 1)

        crop_controls = QHBoxLayout()
        crop_controls.addWidget(QLabel("Ukuran gambar di Excel:"))
        crop_controls.addWidget(self.size_combo)
        crop_controls.addStretch()
        crop_controls.addWidget(self.reset_crop_button)
        crop_controls.addWidget(self.process_button)
        crop_controls.addWidget(self.reprocess_button)
        calibration_layout.addLayout(crop_controls)
        calibration_layout.addWidget(self.progress)
        calibration_layout.addLayout(status_layout)

        calibration_group.setMinimumHeight(560)
        prepare_layout.addWidget(sources, 0)
        prepare_layout.addWidget(calibration_group, 1)
        self.main_tabs.addTab(prepare_page, "1. Import & proses")

        review_group = QGroupBox("Periksa dan lengkapi data")
        review_page = QWidget()
        review_page_layout = QVBoxLayout(review_page)
        review_page_layout.setContentsMargins(8, 8, 8, 8)
        review_layout = QVBoxLayout(review_group)
        self.summary = QLabel("Belum ada data gambar.")
        self.summary.setStyleSheet("font-weight: 600;")
        review_layout.addWidget(self.summary)
        self.table = QTableWidget(0, len(self.table_headers))
        self.table.setHorizontalHeaderLabels(self.table_headers)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(
            QTableWidget.EditTrigger.DoubleClicked |
            QTableWidget.EditTrigger.EditKeyPressed |
            QTableWidget.EditTrigger.AnyKeyPressed
        )
        self.table.verticalHeader().setVisible(False)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setColumnWidth(0, 50)
        self.table.setColumnWidth(1, 105)
        self.table.setColumnWidth(2, 155)
        for column in range(3, 6):
            self.table.setColumnWidth(column, 145)
        for column in range(6, 10):
            self.table.setColumnWidth(column, 125)
        self.table.setColumnWidth(10, 180)
        self.table.itemChanged.connect(self.on_item_changed)
        self.table.cellDoubleClicked.connect(self.preview_original)
        review_layout.addWidget(self.table)
        self.table.cellClicked.connect(self.preview_calibration_row)

        footer = QHBoxLayout()
        self.sort_check = QCheckBox("Urutkan screenshot sesuai urutan Site List")
        self.sort_check.setChecked(True)
        self.duplicate_button = QPushButton("Lihat ID duplikat")
        self.duplicate_button.setEnabled(False)
        self.duplicate_button.clicked.connect(self.show_duplicate_sites)
        self.site_check_button = QPushButton("Cek Site ID & screenshot kosong")
        self.site_check_button.setEnabled(False)
        self.site_check_button.clicked.connect(self.show_site_coverage)
        self.remove_button = QPushButton("Hapus baris terpilih")
        self.remove_button.clicked.connect(self.remove_selected)
        self.export_format_combo = QComboBox()
        self.export_format_combo.addItem(
            "Picture in Cell", EXPORT_FORMAT_PICTURE_IN_CELL
        )
        self.export_format_combo.addItem("Floating", EXPORT_FORMAT_FLOATING)
        self.export_format_combo.setCurrentIndex(
            self.export_format_combo.findData(EXPORT_FORMAT_FLOATING)
        )
        self.export_format_combo.setToolTip(
            "Picture in Cell memerlukan Microsoft Excel yang mendukung fitur "
            "Place in Cell; Floating lebih kompatibel dengan Excel versi lama."
        )
        self.export_button = QPushButton("Ekspor ke template Excel...")
        self.export_button.setObjectName("accentBtn")
        self.export_button.setStyleSheet("font-weight: 700; padding: 7px 14px;")
        self.export_button.clicked.connect(self.export_workbook)
        footer.addWidget(self.sort_check)
        footer.addWidget(self.duplicate_button)
        footer.addWidget(self.site_check_button)
        footer.addStretch()
        footer.addWidget(self.remove_button)
        footer.addWidget(QLabel("Format gambar:"))
        footer.addWidget(self.export_format_combo)
        footer.addWidget(self.export_button)
        review_layout.addLayout(footer)
        review_page_layout.addWidget(review_group, 1)
        self.main_tabs.addTab(review_page, "2. Periksa & lengkapi data")

        self.setCentralWidget(central)
        self._set_busy(False)

    def _apply_theme(self, theme_name: str) -> None:
        self.current_theme_name = theme_name
        application = QApplication.instance()
        if application is not None:
            set_application_theme(theme_name, application)
        self.theme_selector.set_theme(theme_name)
        theme = palette_for(theme_name)
        for canvas in self.calibration_canvases.values():
            canvas.set_theme_colors(theme["card_alt"], theme["muted"])

    def _set_process_status(self, message: str) -> None:
        self.process_status.setText(message)
        self.busy_status_label.setText(message)

    def _show_operation_result(self, success: bool) -> None:
        self.busy_progress.hide()
        if success:
            self.activity_indicator.set_success()
        else:
            self.activity_indicator.set_error()
        self.busy_status_label.setText(self.process_status.text())
        self.busy_banner.show()
        QTimer.singleShot(
            1800,
            lambda: self.busy_banner.hide() if not self.busy else None,
        )

    def reset_all_data(self) -> None:
        if self.busy:
            return
        if not self.items and not self.template_path:
            return
        answer = QMessageBox.question(
            self,
            "Reset data",
            "Hapus semua gambar dan data dari sesi ini, termasuk Site List yang "
            "dipilih? File sumber dan file hasil ekspor tidak akan dihapus.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return
        self.items.clear()
        self.site_map.clear()
        self.site_order.clear()
        self.site_fields.clear()
        self.template_path = None
        self._imported_capture_sources.clear()
        self.duplicate_button.setEnabled(False)
        self.site_check_button.setEnabled(False)
        self.main_tabs.setCurrentIndex(0)
        self.activity_indicator.set_idle()
        self.busy_banner.hide()
        self.table.setRowCount(0)
        self.template_label.setText("Template Excel belum dipilih")
        self.images_label.setText("Belum ada screenshot")
        self.summary.setText("Belum ada data gambar.")
        self.process_status.setText(
            "Site ID diambil dari nama file atau sheet Capture, ZTE, dan FH; "
            "ID dapat dikoreksi manual pada tabel."
        )
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self.sort_check.setChecked(True)
        self.size_combo.setCurrentText("Sedang")
        for canvas in self.calibration_canvases.values():
            canvas.reset_boxes()
            canvas.set_image(None)
        self.calibration_tabs.setCurrentIndex(0)
        self._clipboard_tempdir.cleanup()
        self._embedded_tempdir.cleanup()
        self._clipboard_tempdir = tempfile.TemporaryDirectory(
            prefix="cctv-report-"
        )
        self._embedded_tempdir = tempfile.TemporaryDirectory(
            prefix="cctv-captures-"
        )
        self._clipboard_counter = 0
        self._capture_import_counter = 0

    def choose_template(self) -> None:
        filename, _ = QFileDialog.getOpenFileName(
            self, "Pilih template Site List", str(Path.cwd()),
            "Excel Workbook (*.xlsx)",
        )
        if not filename:
            return
        path = Path(filename)
        try:
            site_map, site_order, site_fields = load_site_data(path)
        except (OSError, ValueError, KeyError) as error:
            QMessageBox.critical(self, "Template tidak dapat dibaca", str(error))
            return
        if not site_map:
            QMessageBox.warning(self, "Site List kosong",
                                "Tidak ada Site ID pada template yang dipilih.")
            return
        self.template_path = path
        self.site_map, self.site_order = site_map, site_order
        self.site_fields = site_fields
        self.site_check_button.setEnabled(True)
        self.template_label.setText(f"{path.name} — {len(site_map)} site")
        for item in self.items:
            if not item.site_name_manual and item.site_id in site_map:
                item.site_name = site_map[item.site_id]
        self.render_table()
        self.update_summary()

    def download_example_template(self, filename: str) -> None:
        available_templates = {
            *REPORT_TEMPLATE_FILES.values(),
            *IMAGE_IMPORT_TEMPLATE_FILES.values(),
        }
        if filename not in available_templates:
            QMessageBox.critical(
                self, "Template tidak tersedia", "Pilihan template tidak valid."
            )
            return

        source = example_template_path(filename)
        if not source.is_file():
            QMessageBox.critical(
                self,
                "Template tidak ditemukan",
                f"File template tidak ditemukan:\n{source}",
            )
            return

        target, _ = QFileDialog.getSaveFileName(
            self,
            "Simpan template CCTV",
            str(Path.home() / "Downloads" / filename),
            "Excel Workbook (*.xlsx)",
        )
        if not target:
            return
        destination = Path(target)
        if destination.suffix.lower() != ".xlsx":
            destination = destination.with_suffix(".xlsx")

        try:
            shutil.copy2(source, destination)
        except OSError as error:
            QMessageBox.critical(
                self, "Gagal menyimpan template", str(error)
            )
            return
        QMessageBox.information(
            self,
            "Template tersimpan",
            f"Template berhasil disimpan:\n{destination}",
        )

    def choose_images(self) -> None:
        filenames, _ = QFileDialog.getOpenFileNames(
            self, "Pilih screenshot HoloSens", str(Path.cwd()),
            "Gambar (*.png *.jpg *.jpeg *.bmp *.webp *.tif *.tiff)",
        )
        if not filenames:
            return
        self._add_images_for_profile(
            [Path(filename) for filename in filenames],
            self.current_calibration_profile(),
        )

    def import_capture_workbook(self) -> None:
        filename, _ = QFileDialog.getOpenFileName(
            self, "Import workbook dengan screenshot tertanam", str(Path.cwd()),
            "Excel Workbook (*.xlsx)",
        )
        if not filename:
            return
        workbook_path = Path(filename)
        resolved_path = workbook_path.resolve()
        if resolved_path in self._imported_capture_sources:
            QMessageBox.information(
                self, APP_NAME, "Workbook ini sudah diimpor pada sesi ini."
            )
            return
        self._capture_import_counter += 1
        output_dir = Path(self._embedded_tempdir.name) / (
            f"import_{self._capture_import_counter:03d}"
        )
        self._set_busy(True)
        self.progress.setRange(0, 0)
        self._set_process_status(
            f"Membaca Site List dan gambar tertanam dari {workbook_path.name}..."
        )
        self.import_worker = CaptureImportWorker(workbook_path, output_dir)
        self.import_worker.progress.connect(self.on_capture_import_progress)
        self.import_worker.completed.connect(self.on_capture_import_completed)
        self.import_worker.finished.connect(self.on_capture_import_worker_finished)
        self.import_worker.start()

    def on_capture_import_progress(self, done: int, total: int, site_id: str) -> None:
        self.progress.setRange(0, total)
        self.progress.setValue(done)
        self._set_process_status(
            f"Mengekstrak screenshot {done}/{total}"
            + (f" — {site_id}" if site_id else "")
        )

    def on_capture_import_completed(self, result: object, error: str) -> None:
        self._set_busy(False)
        if error:
            self._set_process_status("Import workbook gagal.")
            self._show_operation_result(False)
            QMessageBox.critical(self, "Import workbook gagal", error)
            return
        assert isinstance(result, dict)
        workbook_path: Path = result["path"]
        self._imported_capture_sources.add(workbook_path.resolve())
        source_site_map: dict[str, str] = result["site_map"]
        captures: list[ExtractedCapture] = result["captures"]
        if not self.template_path:
            self.template_path = workbook_path
            self.site_map = source_site_map
            self.site_order = result["site_order"]
            self.site_fields = result["site_fields"]
            self.site_check_button.setEnabled(True)
            self.template_label.setText(
                f"{workbook_path.name} — dipakai sebagai workbook laporan"
            )
        added_items: list[ReportItem] = []
        report_crops: dict[str, list[ExtractedCapture]] = {}
        for capture in captures:
            if capture.crop_key and capture.group_key:
                report_crops.setdefault(capture.group_key, []).append(capture)
                continue
            if any(item.path == capture.path for item in self.items):
                continue
            fields = result["site_fields"].get(capture.site_id, {})
            item = ReportItem(
                path=capture.path,
                source_type=capture.source_type,
                site_id=capture.site_id,
                site_name=self.site_map.get(
                    capture.site_id, source_site_map.get(capture.site_id, "")
                ),
                cam1=fields.get("cam1", ""),
                cam2=fields.get("cam2", ""),
                clean=fields.get("clean", ""),
                nvr=fields.get("nvr", ""),
                note=fields.get("note", ""),
            )
            item.manual_fields.update(
                key for key in ("cam1", "cam2", "clean", "note")
                if getattr(item, key)
            )
            item.nvr_manual = bool(item.nvr)
            added_items.append(item)
        for group in report_crops.values():
            first = group[0]
            if any(item.path == first.path for item in self.items):
                continue
            fields = result["site_fields"].get(first.site_id, {})
            item = ReportItem(
                path=first.path,
                source_type=first.source_type,
                site_id=first.site_id,
                site_name=self.site_map.get(
                    first.site_id, source_site_map.get(first.site_id, "")
                ),
                cam1=fields.get("cam1", ""),
                cam2=fields.get("cam2", ""),
                clean=fields.get("clean", ""),
                nvr=fields.get("nvr", ""),
                note=fields.get("note", ""),
                crops={
                    capture.crop_key: capture.path.read_bytes()
                    for capture in group
                    if capture.crop_key
                },
                processed=True,
            )
            item.manual_fields.update(
                key for key in ("cam1", "cam2", "clean", "note")
                if getattr(item, key)
            )
            item.nvr_manual = bool(item.nvr)
            added_items.append(item)
        added = len(added_items)
        self.items.extend(added_items)
        self.items.sort(key=lambda item: natural_key(item.path))
        if added:
            self.refresh_calibration_images()
        self.images_label.setText(
            f"{len(self.items)} screenshot dipilih ({added} dari workbook)"
        )
        crop_count = sum(capture.crop_key is not None for capture in captures)
        full_screenshot_count = len(captures) - crop_count
        self.progress.setRange(0, 100)
        self.progress.setValue(100)
        self.render_table()
        self.update_summary()
        self._set_process_status(
            f"✓ Selesai: {added} gambar/laporan diimpor dari "
            f"{workbook_path.name} ({full_screenshot_count} screenshot, "
            f"{crop_count} crop laporan)."
        )
        self._show_operation_result(True)

    def on_capture_import_worker_finished(self) -> None:
        self.import_worker = None

    def paste_clipboard(self) -> None:
        image = QApplication.clipboard().image()
        if image.isNull():
            QMessageBox.information(
                self, APP_NAME,
                "Clipboard belum berisi gambar. Salin screenshot, lalu tekan Ctrl+V.",
            )
            return
        self._clipboard_counter += 1
        path = Path(self._clipboard_tempdir.name) / (
            f"clipboard_{self._clipboard_counter:04d}.png"
        )
        if not image.save(str(path), "PNG"):
            QMessageBox.warning(self, APP_NAME,
                                "Gambar dari clipboard tidak dapat disimpan.")
            return
        self._add_images_for_profile([path], self.current_calibration_profile())

    def eventFilter(self, watched: Any, event: Any) -> bool:
        if (event.type() == QEvent.Type.KeyPress
                and event.key() == Qt.Key.Key_V
                and not self.busy
                and event.modifiers() & Qt.KeyboardModifier.ControlModifier
                and QApplication.clipboard().mimeData().hasImage()):
            self.paste_clipboard()
            return True
        return super().eventFilter(watched, event)

    def _add_images(self, paths: list[Path]) -> None:
        self._add_images_for_profile(paths, self.current_calibration_profile())

    def _add_images_for_profile(self, paths: list[Path],
                                source_type: str) -> None:
        known = {str(item.path.resolve()).casefold() for item in self.items}
        added = []
        for path in paths:
            resolved_path = str(path.resolve()).casefold()
            if resolved_path in known:
                continue
            item = ReportItem(path=path, source_type=source_type,
                              site_id=site_id_from_filename(path))
            item.site_name = self.site_map.get(item.site_id, "")
            self.items.append(item)
            added.append(item)
            known.add(resolved_path)
        self.items.sort(key=lambda item: natural_key(item.path))
        if added:
            self.refresh_calibration_images()
        self.images_label.setText(f"{len(self.items)} screenshot dipilih")
        self.render_table()
        self.update_summary()

    def current_calibration_profile(self) -> str:
        index = self.calibration_tabs.currentIndex()
        return PROFILE_KEYS[index]

    def reset_current_calibration(self) -> None:
        self.calibration_canvases[
            self.current_calibration_profile()
        ].reset_boxes()

    def refresh_calibration_images(self) -> None:
        for profile, canvas in self.calibration_canvases.items():
            item = next((record for record in self.items
                         if record.source_type == profile), None)
            canvas.set_image(item.path if item else None)

    def preview_calibration_row(self, row: int, column: int) -> None:
        del column
        if not 0 <= row < len(self.items):
            return
        profile = self.items[row].source_type
        self.calibration_tabs.setCurrentIndex(PROFILE_KEYS.index(profile))
        self.calibration_canvases[profile].set_image(self.items[row].path)

    def process_images(self, reprocess: bool = False) -> None:
        if self.busy:
            return
        pending = [item for item in self.items if reprocess or not item.processed]
        if not pending:
            QMessageBox.information(self, APP_NAME, "Tidak ada screenshot yang perlu diproses.")
            return
        panel_width, camera_width, quality = IMAGE_SIZES[
            self.size_combo.currentText()
        ]
        self._set_busy(True)
        self.progress.setRange(0, len(pending))
        self.progress.setValue(0)
        self.worker = ProcessingWorker(
            pending,
            {profile: canvas.normalized_boxes()
             for profile, canvas in self.calibration_canvases.items()},
            panel_width,
            camera_width,
            quality,
        )
        self.worker.progress.connect(self.on_progress)
        self.worker.item_finished.connect(
            lambda index, result, error: self.on_processed_item(
                pending[index], result, error
            )
        )
        self.worker.finished.connect(self.on_processing_finished)
        self.worker.start()

    def reprocess_images(self) -> None:
        self.process_images(reprocess=True)

    def on_processed_item(self, item: ReportItem, result: object, error: str) -> None:
        if error:
            item.processed = False
            item.crops.clear()
            item.note = f"Proses gagal: {error}"
            return
        assert isinstance(result, dict)
        item.crops = result["crops"]
        for key, value in result["statuses"].items():
            if key not in item.manual_fields:
                setattr(item, key, value)
        if not item.nvr_manual:
            item.nvr = derive_nvr(item.cam1, item.cam2)
        if "note" not in item.manual_fields:
            item.note = result["note"]
        item.processed = True

    def on_progress(self, done: int, total: int, filename: str) -> None:
        self.progress.setValue(done)
        self._set_process_status(f"Memproses {done}/{total}: {filename}")

    def on_processing_finished(self) -> None:
        failed = sum(not item.processed for item in self.items)
        completed = sum(item.processed for item in self.items)
        self._set_busy(False)
        self.render_table()
        self.update_summary()
        self._set_process_status(
            f"Selesai: {completed} gambar berhasil"
            + (f", {failed} gagal. Periksa nama/file gambar lalu coba lagi." if failed else ".")
        )
        if failed:
            self._show_operation_result(False)
        else:
            self._show_operation_result(True)

    def render_table(self) -> None:
        self._loading_table = True
        self.table.setRowCount(len(self.items))
        for row, record in enumerate(self.items):
            number = QTableWidgetItem(str(row + 1))
            number.setFlags(number.flags() & ~Qt.ItemFlag.ItemIsEditable)
            self.table.setItem(row, 0, number)
            values = (record.site_id, record.site_name)
            for offset, value in enumerate(values, start=1):
                self.table.setItem(row, offset, QTableWidgetItem(value))
            for offset, key in enumerate(("panel", "cam1", "cam2"), start=3):
                self.table.setCellWidget(row, offset, self._thumbnail(record.crops.get(key)))
            for offset, key, choices in (
                (6, "cam1", CAMERA_STATUSES),
                (7, "cam2", CAMERA_STATUSES),
                (8, "clean", CLEAN_STATUSES),
                (9, "nvr", CAMERA_STATUSES),
            ):
                combo = QComboBox()
                combo.addItem("")
                combo.addItems(choices)
                combo.setCurrentText(getattr(record, key))
                combo.currentTextChanged.connect(
                    lambda value, idx=row, field_name=key:
                    self.on_status_changed(idx, field_name, value)
                )
                self.table.setCellWidget(row, offset, combo)
            note_item = QTableWidgetItem(record.note)
            self.table.setItem(row, 10, note_item)
            self.table.setRowHeight(row, 86)
        self._loading_table = False
        self.update_summary()

    @staticmethod
    def _thumbnail(data: bytes | None) -> QLabel:
        label = QLabel("Belum diproses" if not data else "")
        label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        label.setMinimumSize(130, 78)
        if data:
            image = QImage.fromData(data)
            if not image.isNull():
                label.setPixmap(QPixmap.fromImage(image).scaled(
                    135, 78, Qt.AspectRatioMode.KeepAspectRatio,
                    Qt.TransformationMode.SmoothTransformation,
                ))
        return label

    def on_item_changed(self, cell: QTableWidgetItem) -> None:
        if self._loading_table or not 0 <= cell.row() < len(self.items):
            return
        item = self.items[cell.row()]
        if cell.column() == 1:
            item.site_id = re.sub(r"\s+", "", cell.text().upper())
            cell.setText(item.site_id)
            if not item.site_name_manual:
                item.site_name = self.site_map.get(item.site_id, "")
                self._loading_table = True
                self.table.item(cell.row(), 2).setText(item.site_name)
                self._loading_table = False
        elif cell.column() == 2:
            item.site_name = cell.text()
            item.site_name_manual = True
        elif cell.column() == 10:
            item.note = cell.text()
            item.manual_fields.add("note")
        self.update_summary()

    def on_status_changed(self, row: int, field_name: str, value: str) -> None:
        if self._loading_table or not 0 <= row < len(self.items):
            return
        item = self.items[row]
        setattr(item, field_name, value)
        item.manual_fields.add(field_name)
        if field_name in ("cam1", "cam2"):
            if not item.nvr_manual:
                item.nvr = derive_nvr(item.cam1, item.cam2)
                nvr_combo = self.table.cellWidget(row, 9)
                if isinstance(nvr_combo, QComboBox):
                    nvr_combo.blockSignals(True)
                    nvr_combo.setCurrentText(item.nvr)
                    nvr_combo.blockSignals(False)
        elif field_name == "nvr":
            item.nvr_manual = True
        self.update_summary()

    def update_summary(self) -> None:
        if not self.items:
            self.summary.setText("Belum ada data gambar.")
            self.duplicate_button.setEnabled(False)
            return
        ids = [
            item.site_id.strip().upper()
            for item in self.items if item.site_id.strip()
        ]
        duplicates = len(ids) - len(set(ids))
        self.duplicate_button.setEnabled(not self.busy and duplicates > 0)
        missing_id = sum(not item.site_id for item in self.items)
        unknown = sum(bool(item.site_id and self.site_map and
                           item.site_id not in self.site_map) for item in self.items)
        covered = len(set(ids) & self.site_map.keys()) if self.site_map else 0
        text = (f"{len(self.items)} screenshot   |   {missing_id} ID kosong   |   "
                f"{duplicates} ID duplikat   |   {unknown} ID tidak ada di Site List")
        if self.site_map:
            text += f"   |   {covered}/{len(self.site_map)} site tercakup"
        profiles = Counter(item.source_type for item in self.items)
        text += (
            f"   |   {profiles['ZTE']} ZTE / {profiles['FH']} FH / "
            f"{profiles['P11']} Paket 11"
        )
        self.summary.setText(text)

    def _duplicate_item_indices(self) -> list[int]:
        counts = Counter(
            item.site_id.strip().upper()
            for item in self.items if item.site_id.strip()
        )
        return [
            index for index, item in enumerate(self.items)
            if item.site_id.strip()
            and counts[item.site_id.strip().upper()] > 1
        ]

    def show_duplicate_sites(self) -> None:
        if self.busy:
            return
        dialog = QDialog(self)
        dialog.setWindowTitle("Site ID duplikat")
        dialog.setMinimumSize(760, 360)
        layout = QVBoxLayout(dialog)
        description = QLabel(
            "Pilih salah satu baris duplikat untuk menghapus screenshot tersebut. "
            "Baris yang tersisa tidak akan dihapus."
        )
        description.setWordWrap(True)
        layout.addWidget(description)

        duplicates_table = QTableWidget(0, 5)
        duplicates_table.setHorizontalHeaderLabels(
            ("Site ID", "Urutan", "Site Name", "Tipe", "Nama file")
        )
        duplicates_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        duplicates_table.setSelectionMode(
            QTableWidget.SelectionMode.SingleSelection
        )
        duplicates_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        duplicates_table.verticalHeader().setVisible(False)
        duplicates_table.horizontalHeader().setStretchLastSection(True)
        duplicates_table.setColumnWidth(0, 105)
        duplicates_table.setColumnWidth(1, 80)
        duplicates_table.setColumnWidth(2, 180)
        duplicates_table.setColumnWidth(3, 150)
        layout.addWidget(duplicates_table)

        status = QLabel()
        layout.addWidget(status)
        buttons = QHBoxLayout()
        delete_button = QPushButton("Hapus screenshot terpilih")
        close_button = QPushButton("Tutup")
        close_button.clicked.connect(dialog.accept)
        buttons.addStretch()
        buttons.addWidget(delete_button)
        buttons.addWidget(close_button)
        layout.addLayout(buttons)

        def refresh_duplicates() -> None:
            duplicate_indices = self._duplicate_item_indices()
            occurrences: Counter[str] = Counter()
            duplicates_table.setRowCount(len(duplicate_indices))
            for row, index in enumerate(duplicate_indices):
                record = self.items[index]
                site_id = record.site_id.strip().upper()
                occurrences[site_id] += 1
                values = (
                    site_id,
                    str(occurrences[site_id]),
                    record.site_name,
                    PROFILE_LABELS.get(record.source_type, record.source_type),
                    record.path.name,
                )
                for column, value in enumerate(values):
                    cell = QTableWidgetItem(value)
                    if column == 0:
                        cell.setData(Qt.ItemDataRole.UserRole, index)
                    duplicates_table.setItem(row, column, cell)
            has_duplicates = bool(duplicate_indices)
            delete_button.setEnabled(has_duplicates)
            status.setText(
                f"{len(duplicate_indices)} screenshot pada ID duplikat."
                if has_duplicates else "Tidak ada lagi Site ID duplikat."
            )

        def delete_selected_duplicate() -> None:
            selected_rows = duplicates_table.selectionModel().selectedRows()
            if not selected_rows:
                QMessageBox.warning(
                    dialog, "Pilih screenshot",
                    "Pilih salah satu baris duplikat yang ingin dihapus.",
                )
                return
            selected_row = selected_rows[0].row()
            site_id_cell = duplicates_table.item(selected_row, 0)
            if site_id_cell is None:
                return
            item_index = site_id_cell.data(Qt.ItemDataRole.UserRole)
            if not isinstance(item_index, int) or not 0 <= item_index < len(self.items):
                return
            record = self.items[item_index]
            answer = QMessageBox.question(
                dialog,
                "Konfirmasi hapus",
                f"Hapus screenshot {record.path.name} dengan Site ID "
                f"{record.site_id} dari laporan ini?\n\nFile sumber tidak dihapus.",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.No,
            )
            if answer != QMessageBox.StandardButton.Yes:
                return
            del self.items[item_index]
            self.images_label.setText(f"{len(self.items)} screenshot dipilih")
            self.refresh_calibration_images()
            self.render_table()
            self.update_summary()
            refresh_duplicates()

        delete_button.clicked.connect(delete_selected_duplicate)
        refresh_duplicates()
        dialog.exec()

    def show_site_coverage(self) -> None:
        if self.busy:
            return
        if not self.template_path:
            QMessageBox.information(
                self, APP_NAME,
                "Pilih template atau import workbook yang memiliki Site List "
                "terlebih dahulu.",
            )
            return
        try:
            duplicate_template_sites = find_duplicate_site_rows(
                self.template_path
            )
        except (OSError, ValueError, KeyError) as error:
            QMessageBox.critical(
                self, "Gagal memeriksa Site ID", str(error)
            )
            return

        covered_ids = {
            item.site_id.strip().upper()
            for item in self.items
            if item.site_id.strip() and (item.path.is_file() or item.crops)
        }
        missing_sites = [
            (site_id, site_name)
            for site_id, site_name in self.site_map.items()
            if site_id not in covered_ids
        ]
        blank_id_count = sum(
            not item.site_id.strip()
            and (item.path.is_file() or item.crops)
            for item in self.items
        )
        duplicate_capture_sites = [
            (site_id, sum(
                item.site_id.strip().upper() == site_id for item in self.items
            ))
            for site_id in dict.fromkeys(
                item.site_id.strip().upper()
                for item in self.items if item.site_id.strip()
            )
            if sum(
                item.site_id.strip().upper() == site_id for item in self.items
            ) > 1
        ]

        dialog = QDialog(self)
        dialog.setWindowTitle("Pemeriksaan Site ID dan screenshot")
        dialog.setMinimumSize(780, 460)
        layout = QVBoxLayout(dialog)
        summary = QLabel(
            f"{len(self.site_map)} Site ID di template | "
            f"{len(covered_ids)} site memiliki screenshot | "
            f"{len(missing_sites)} site belum memiliki screenshot"
            + (f" | {blank_id_count} screenshot tanpa Site ID" if blank_id_count else "")
        )
        summary.setWordWrap(True)
        layout.addWidget(summary)

        tabs = QTabWidget()
        duplicate_table = QTableWidget(0, 4)
        duplicate_table.setHorizontalHeaderLabels(
            ("Sumber", "Site ID", "Site Name / File", "Jumlah")
        )
        duplicate_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        duplicate_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        duplicate_table.verticalHeader().setVisible(False)
        duplicate_table.horizontalHeader().setStretchLastSection(True)
        duplicate_table.setColumnWidth(0, 170)
        duplicate_table.setColumnWidth(1, 120)
        duplicate_table.setColumnWidth(3, 75)
        duplicate_rows: list[tuple[str, str, str, int]] = [
            ("Template Site List", site_id, site_name, count)
            for site_id, site_name, count in duplicate_template_sites
        ]
        duplicate_rows.extend(
            (
                "Screenshot di sesi",
                site_id,
                ", ".join(
                    item.path.name for item in self.items
                    if item.site_id.strip().upper() == site_id
                ),
                count,
            )
            for site_id, count in duplicate_capture_sites
        )
        duplicate_table.setRowCount(len(duplicate_rows))
        for row, values in enumerate(duplicate_rows):
            for column, value in enumerate(values):
                duplicate_table.setItem(
                    row, column, QTableWidgetItem(str(value))
                )
        duplicate_label = QLabel(
            f"{len(duplicate_template_sites)} ID duplikat di template; "
            f"{len(duplicate_capture_sites)} ID memiliki beberapa screenshot."
        )
        duplicate_label.setWordWrap(True)
        duplicate_page = QWidget()
        duplicate_layout = QVBoxLayout(duplicate_page)
        duplicate_layout.addWidget(duplicate_label)
        duplicate_layout.addWidget(duplicate_table)
        tabs.addTab(duplicate_page, "Site ID duplikat")

        missing_table = QTableWidget(0, 2)
        missing_table.setHorizontalHeaderLabels(("Site ID", "Site Name"))
        missing_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        missing_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        missing_table.verticalHeader().setVisible(False)
        missing_table.horizontalHeader().setStretchLastSection(True)
        missing_table.setColumnWidth(0, 150)
        missing_table.setRowCount(len(missing_sites))
        for row, (site_id, site_name) in enumerate(missing_sites):
            missing_table.setItem(row, 0, QTableWidgetItem(site_id))
            missing_table.setItem(row, 1, QTableWidgetItem(site_name))
        missing_page = QWidget()
        missing_layout = QVBoxLayout(missing_page)
        missing_layout.addWidget(QLabel(
            "Site dari template yang belum memiliki screenshot di sesi ini."
        ))
        missing_layout.addWidget(missing_table)
        tabs.addTab(missing_page, f"Belum ada screenshot ({len(missing_sites)})")
        layout.addWidget(tabs)

        close_button = QPushButton("Tutup")
        close_button.clicked.connect(dialog.accept)
        buttons = QHBoxLayout()
        buttons.addStretch()
        buttons.addWidget(close_button)
        layout.addLayout(buttons)
        dialog.exec()

    def preview_original(self, row: int, column: int) -> None:
        if column not in (3, 4, 5) or not 0 <= row < len(self.items):
            return
        path = self.items[row].path
        try:
            image = QImage(str(path))
        except (OSError, ValueError):
            return
        if image.isNull():
            return
        dialog = QMessageBox(self)
        dialog.setWindowTitle(path.name)
        dialog.setText(f"Screenshot asli: {path}")
        dialog.setIconPixmap(QPixmap.fromImage(image).scaled(
            900, 650, Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation,
        ))
        dialog.exec()

    def remove_selected(self) -> None:
        selected = sorted({index.row() for index in self.table.selectionModel().selectedRows()},
                          reverse=True)
        for row in selected:
            del self.items[row]
        self.images_label.setText(f"{len(self.items)} screenshot dipilih")
        self.refresh_calibration_images()
        self.render_table()
        self.update_summary()

    def export_workbook(self) -> None:
        if not self.template_path:
            QMessageBox.warning(self, APP_NAME, "Pilih template Excel terlebih dahulu.")
            return
        if self.busy:
            return
        records = [item for item in self.items if item.processed]
        if not records:
            QMessageBox.warning(self, APP_NAME, "Proses setidaknya satu screenshot terlebih dahulu.")
            return
        if self.sort_check.isChecked() and self.site_order:
            records.sort(key=lambda item: (
                self.site_order.get(item.site_id, 10**9),
                natural_key(item.path),
            ))
        default_name = f"{self.template_path.stem}_report.xlsx"
        filename, _ = QFileDialog.getSaveFileName(
            self, "Simpan laporan CCTV", str(self.template_path.with_name(default_name)),
            "Excel Workbook (*.xlsx)",
        )
        if not filename:
            return
        output_path = Path(filename)
        if output_path.suffix.casefold() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")
        if output_path.resolve() == self.template_path.resolve():
            QMessageBox.warning(self, APP_NAME,
                                "Pilih nama file output yang berbeda dari template sumber.")
            return
        selected_format = str(self.export_format_combo.currentData() or "")
        if selected_format == EXPORT_FORMAT_PICTURE_IN_CELL:
            image_placement: ExportFormat = EXPORT_FORMAT_PICTURE_IN_CELL
        elif selected_format == EXPORT_FORMAT_FLOATING:
            image_placement = EXPORT_FORMAT_FLOATING
        else:
            QMessageBox.warning(self, APP_NAME, "Pilih format gambar ekspor.")
            return
        self._set_busy(True)
        self.progress.setRange(0, 0)
        self._set_process_status("Memulai ekspor ke template Excel...")
        self.export_worker = ExportWorker(
            self, output_path, records, image_placement
        )
        self.export_worker.progress.connect(self.on_export_progress)
        self.export_worker.completed.connect(self.on_export_completed)
        self.export_worker.finished.connect(self.on_export_worker_finished)
        self._last_export_progress_at = time.monotonic()
        self._last_export_progress_message = self.process_status.text()
        self.export_watchdog.start()
        self.export_worker.start()

    def on_export_progress(self, done: int, total: int, message: str) -> None:
        self._last_export_progress_at = time.monotonic()
        self._last_export_progress_message = message
        if total > 0:
            self.busy_progress.setRange(0, total)
            self.busy_progress.setValue(min(done, total))
        else:
            self.busy_progress.setRange(0, 0)
        self._set_process_status(message)

    def _update_export_watchdog(self) -> None:
        if not self.busy or not self.export_worker or not self.export_worker.isRunning():
            self.export_watchdog.stop()
            return
        if not self._last_export_progress_at:
            return
        stalled_for = int(time.monotonic() - self._last_export_progress_at)
        if stalled_for >= 60:
            self.busy_status_label.setText(
                f"{self._last_export_progress_message}\n"
                f"Belum ada pembaruan selama {stalled_for // 60} menit "
                f"{stalled_for % 60} detik. Periksa apakah Excel masih merespons."
            )

    def on_export_completed(self, result: object, error: str) -> None:
        self.export_watchdog.stop()
        self._last_export_progress_at = 0.0
        self._set_busy(False)
        if error:
            self.progress.setRange(0, 100)
            self.progress.setValue(0)
            self._set_process_status("Ekspor gagal.")
            self._show_operation_result(False)
            QMessageBox.critical(self, "Ekspor gagal", error)
            return
        assert isinstance(result, tuple) and len(result) == 3
        output_path, counts, image_placement = result
        placed, floating, warning = counts
        self.progress.setRange(0, 100)
        self.progress.setValue(100)
        self._set_process_status("✓ Ekspor ke template Excel selesai.")
        self._show_operation_result(True)
        message = f"Laporan berhasil disimpan:\n{output_path}"
        if image_placement == EXPORT_FORMAT_FLOATING:
            message += f"\n\nFormat gambar: Floating ({floating} gambar)."
        else:
            if placed:
                message += f"\n\nPicture in Cell: {placed} gambar."
            if floating:
                message += (
                    f"\nGambar floating yang dipaskan ke sel: {floating}."
                )
        if warning:
            message += f"\n\n{warning}"
            QMessageBox.warning(self, "Ekspor selesai dengan fallback", message)
        else:
            QMessageBox.information(self, "Ekspor selesai", message)

    def on_export_worker_finished(self) -> None:
        self.export_worker = None

    def _write_report(self, output_path: Path,
                      records: list[ReportItem]) -> int:
        assert self.template_path is not None
        workbook = load_workbook(self.template_path)
        worksheet, header_row, id_column, name_column = find_site_columns(workbook)
        headers = {
            re.sub(r"\s+", " ", str(worksheet.cell(header_row, column).value or ""))
            .strip().casefold(): column
            for column in range(1, worksheet.max_column + 1)
        }
        columns: dict[str, int] = {"site_id": id_column, "site_name": name_column}
        columns["nvr"] = headers.get("nvr", 0)
        columns["cam1"] = (
            find_camera_status_column(headers, "camera 1", "kabinet") or 0
        )
        columns["cam2"] = (
            find_camera_status_column(headers, "camera 2", "pintu pagar") or 0
        )
        columns["note"] = headers.get("keterangan", 0)
        columns["no"] = headers.get("no", 0)
        columns["panel_image"] = headers.get("sensor monitoring", 0)
        columns["cam1_image"] = headers.get("camera 1", 0)
        columns["cam2_image"] = headers.get("camera 2", 0)
        has_report_images = all(
            columns[key] for key in ("panel_image", "cam1_image", "cam2_image")
        )
        if not all(columns[key] for key in ("nvr", "cam1", "cam2", "note")):
            workbook.close()
            raise ValueError(
                "Workbook harus memiliki kolom NVR, Camera 1, Camera 2, dan Keterangan."
            )

        columns["clean"] = headers.get("kebersihan", 0)
        if not columns["clean"]:
            columns["clean"] = worksheet.max_column + 1
            style_column = columns["note"]
            for row in range(1, worksheet.max_row + 1):
                source_cell = worksheet.cell(row, style_column)
                target_cell = worksheet.cell(row, columns["clean"])
                if source_cell.has_style:
                    target_cell._style = copy(source_cell._style)
                target_cell.alignment = copy(source_cell.alignment)
            worksheet.cell(header_row, columns["clean"]).value = "Kebersihan"
            worksheet.column_dimensions[
                get_column_letter(columns["clean"])
            ].width = 16

        id_rows: dict[str, list[int]] = {}
        record_rows: dict[int, int] = {}
        last_data_row = header_row
        for row in range(header_row + 1, worksheet.max_row + 1):
            site_id = str(worksheet.cell(row, id_column).value or "").strip().upper()
            if not site_id:
                continue
            id_rows.setdefault(site_id, []).append(row)
            last_data_row = row

        camera_values = {
            str(worksheet.cell(row, columns[key]).value or "").strip().upper()
            for key in ("cam1", "cam2")
            for row in range(header_row + 1, last_data_row + 1)
        }
        camera_encoding = "ok_nok" if camera_values & {"OK", "NOK"} else "standard"
        nvr_values = {
            str(worksheet.cell(row, columns["nvr"]).value or "").strip()
            for row in range(header_row + 1, last_data_row + 1)
        }
        uppercase_status = any(
            value in {"ONLINE", "OFFLINE", "NOT DETECTED"}
            for value in nvr_values
        )
        extra_rows: list[tuple[ReportItem, int, int | None]] = []
        used: dict[str, int] = {}
        for record in records:
            matches = id_rows.get(record.site_id, [])
            occurrence = used.get(record.site_id, 0)
            used[record.site_id] = occurrence + 1
            if occurrence < len(matches):
                row = matches[occurrence]
                self._write_record(worksheet, row, record, columns, header_row,
                                   camera_encoding, uppercase_status)
            else:
                last_data_row += 1
                prototype_row = next(iter(id_rows.values()))[0] if id_rows else header_row + 1
                self._copy_row_style(worksheet, prototype_row, last_data_row)
                extra_rows.append((record, last_data_row, prototype_row))
                self._write_record(worksheet, last_data_row, record, columns,
                                   header_row, camera_encoding, uppercase_status)
                row = last_data_row
            record_rows[id(record)] = row

        for record, row, prototype_row in extra_rows:
            for column in (
                headers.get("konsorsium pembangunan", 0),
                headers.get("paket penyedia mss", 0),
            ):
                if not column:
                    continue
                worksheet.cell(row, column).value = worksheet.cell(prototype_row, column).value
        for row in range(header_row + 1, last_data_row + 1):
            if columns["no"] and worksheet.cell(row, id_column).value:
                worksheet.cell(row, columns["no"]).value = row - header_row
        last_column = max(worksheet.max_column, max(columns.values()))
        worksheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(last_column)}{last_data_row}"
            if last_data_row > header_row else None
        )
        image_count = 0
        if has_report_images:
            replaced_cells = {
                (columns[column_key] - 1, record_rows[id(record)] - 1)
                for record in records
                for column_key in ("panel_image", "cam1_image", "cam2_image")
            }
            worksheet._images = [
                image for image in worksheet._images
                if (
                    getattr(getattr(image.anchor, "_from", None), "col", None),
                    getattr(getattr(image.anchor, "_from", None), "row", None),
                ) not in replaced_cells
            ]
            for key, width in (("panel_image", 30), ("cam1_image", 32),
                               ("cam2_image", 32)):
                worksheet.column_dimensions[
                    get_column_letter(columns[key])
                ].width = width
            worksheet.column_dimensions[get_column_letter(id_column)].width = 16
            worksheet.column_dimensions[get_column_letter(name_column)].width = 25
            worksheet.column_dimensions[
                get_column_letter(columns["note"])
            ].width = 30
            self._update_status_validation(
                worksheet, header_row, last_data_row, columns, has_report_images
            )
            for record in records:
                target_row = record_rows[id(record)]
                for key, column_key in (
                    ("panel", "panel_image"),
                    ("cam1", "cam1_image"),
                    ("cam2", "cam2_image"),
                ):
                    data = record.crops.get(key)
                    if not data:
                        continue
                    excel_image = SpreadsheetImage(io.BytesIO(data))
                    max_width = 204 if key == "panel" else 218
                    scale = min(
                        max_width / excel_image.width,
                        160 / excel_image.height,
                        1,
                    )
                    excel_image.width = round(excel_image.width * scale)
                    excel_image.height = round(excel_image.height * scale)
                    worksheet.add_image(
                        excel_image,
                        f"{get_column_letter(columns[column_key])}{target_row}",
                    )
                    image_count += 1
                worksheet.row_dimensions[target_row].height = 130.05
        else:
            self._update_status_validation(
                worksheet, header_row, last_data_row, columns, has_report_images
            )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        workbook.close()
        return image_count

    @staticmethod
    def _place_images_in_cells(
        output_path: Path,
        progress_callback: Any | None = None,
    ) -> tuple[int, int, str]:
        try:
            with zipfile.ZipFile(output_path) as archive:
                drawing_files = [
                    name for name in archive.namelist()
                    if re.fullmatch(r"xl/drawings/drawing\d+\.xml", name)
                ]
                image_count = sum(
                    len(ET.fromstring(archive.read(name)).findall(
                        ".//xdr:pic", OOXML_NS
                    ))
                    for name in drawing_files
                )
        except (OSError, zipfile.BadZipFile, ET.ParseError) as error:
            return (
                0,
                0,
                f"Gambar dalam workbook tidak dapat dihitung ({error}); "
                "ekspor tetap tersimpan.",
            )

        try:
            import win32com.client
        except ImportError as error:
            if progress_callback:
                progress_callback(
                    0, image_count,
                    f"Picture in Cell tidak tersedia: {image_count} gambar "
                    "belum dikonversi.",
                )
            return (
                0,
                image_count,
                f"Microsoft Excel automation tidak tersedia ({error}). "
                "Gambar dibiarkan floating sesuai fallback.",
            )

        application = None
        workbook = None
        placed_count = 0
        warning = ""
        command_ids = (
            "PicturePlaceInCell",
            "PlaceInCell",
            "PictureInCell",
            "PlacePictureInCell",
            "PicturesPlaceInCell",
            "ObjectPlaceInCell",
            "InsertPictureInCell",
            "PictureFormatPlaceInCell",
            "PlacePictureInCellGallery",
        )
        try:
            application = win32com.client.DispatchEx("Excel.Application")
            application.Visible = False
            application.DisplayAlerts = False
            application.AskToUpdateLinks = False
            application.EnableEvents = False
            application.AutomationSecurity = 3
            workbook = application.Workbooks.Open(
                str(output_path.resolve()), UpdateLinks=0, ReadOnly=False
            )

            command_id = ""
            for candidate in command_ids:
                try:
                    application.CommandBars.GetEnabledMso(candidate)
                except Exception:
                    continue
                command_id = candidate
                break

            picture_shapes: list[tuple[Any, str]] = []
            for sheet_index in range(1, workbook.Worksheets.Count + 1):
                worksheet = workbook.Worksheets.Item(sheet_index)
                for shape_index in range(1, worksheet.Shapes.Count + 1):
                    shape = worksheet.Shapes.Item(shape_index)
                    if int(shape.Type) in (11, 13):
                        picture_shapes.append((worksheet, str(shape.Name)))
            image_count = max(image_count, len(picture_shapes))
            started_at = time.monotonic()
            failed_shapes: list[tuple[str, str]] = []

            def report_progress(attempted: int) -> None:
                if not progress_callback:
                    return
                elapsed = max(0.001, time.monotonic() - started_at)
                rate = attempted / elapsed
                remaining = image_count - attempted
                eta = ""
                if attempted >= 10 and rate > 0:
                    eta_seconds = round(remaining / rate)
                    eta = (
                        f" | sisa ~{eta_seconds // 60} mnt "
                        f"{eta_seconds % 60} dtk"
                    )
                progress_callback(
                    attempted,
                    image_count,
                    f"Picture in Cell {attempted}/{image_count} | "
                    f"{placed_count} berhasil, {len(failed_shapes)} gagal"
                    f"{eta}",
                )

            report_progress(0)

            if not command_id:
                warning = (
                    "Excel tidak mengenali perintah Place in Cell. "
                    "Gambar dibiarkan floating sesuai fallback."
                )
                failed_shapes.extend(
                    (shape_name, "Place in Cell tidak tersedia")
                    for _, shape_name in picture_shapes
                )
                if progress_callback:
                    progress_callback(
                        0,
                        image_count,
                        f"Picture in Cell 0/{image_count} | 0 berhasil, "
                        f"{len(failed_shapes)} gagal",
                    )
            else:
                for attempted, (worksheet, shape_name) in enumerate(
                    picture_shapes, start=1
                ):
                    try:
                        worksheet.Activate()
                        shapes = worksheet.Shapes
                        count_before = int(shapes.Count)
                        shapes.Item(shape_name).Select()
                        application.CommandBars.ExecuteMso(command_id)
                        try:
                            shapes.Item(shape_name)
                            still_present = True
                        except Exception:
                            still_present = False
                        count_after = int(shapes.Count)
                        if still_present or count_after >= count_before:
                            failed_shapes.append(
                                (shape_name, "Objek gambar masih floating")
                            )
                        else:
                            placed_count += 1
                    except Exception as error:
                        failed_shapes.append((shape_name, str(error)))
                    report_progress(attempted)

                if placed_count < image_count:
                    failure_details = "; ".join(
                        f"{name}: {reason}"
                        for name, reason in failed_shapes[:5]
                    )
                    omitted_failures = len(failed_shapes) - 5
                    warning = (
                        "Sebagian gambar tidak dapat diubah ke Picture in Cell "
                        f"oleh Excel ({placed_count}/{image_count} berhasil). "
                        "Gambar yang gagal tetap floating di selnya."
                    )
                    if failure_details:
                        warning += f" Contoh: {failure_details}."
                    if omitted_failures > 0:
                        warning += f" Dan {omitted_failures} kegagalan lainnya."
            if progress_callback:
                progress_callback(
                    image_count,
                    image_count,
                    "Menyimpan perubahan Picture in Cell ke file Excel...",
                )
            workbook.Save()
        except Exception as error:
            warning = (
                f"Picture in Cell tidak dapat diproses melalui Microsoft Excel "
                f"({error}). Gambar yang belum terkonversi tetap floating."
            )
        finally:
            if workbook is not None:
                try:
                    workbook.Close(SaveChanges=True)
                except Exception as error:
                    warning = (
                        f"{warning} Workbook Excel gagal ditutup dengan bersih: "
                        f"{error}"
                    ).strip()
            if application is not None:
                try:
                    application.Quit()
                except Exception as error:
                    warning = (
                        f"{warning} Microsoft Excel gagal ditutup dengan bersih: "
                        f"{error}"
                    ).strip()

        return placed_count, max(0, image_count - placed_count), warning

    @staticmethod
    def _update_status_validation(worksheet: Any, header_row: int,
                                  last_data_row: int, columns: dict[str, int],
                                  has_report_images: bool) -> None:
        if has_report_images:
            status_columns = [columns[key] for key in ("nvr", "cam1", "cam2", "clean")]
            min_column, max_column = min(status_columns), max(status_columns)
            validation = next(
                (item for item in worksheet.data_validations.dataValidation
                 if item.type == "list" and any(
                     area.min_col <= min_column and area.max_col >= max_column
                     for area in item.sqref.ranges
                 )),
                None,
            )
            formula = ('"Online,Offline,Not Detected,Online With Note,'
                       'Bersih,Kurang Bersih,Tidak Bersih"')
            if validation:
                validation.formula1 = formula
                validated_through = max(
                    (area.max_row for area in validation.sqref.ranges),
                    default=header_row,
                )
                if last_data_row > validated_through:
                    validation.add(
                        f"{get_column_letter(min_column)}{validated_through + 1}:"
                        f"{get_column_letter(max_column)}{last_data_row}"
                    )
            else:
                validation = DataValidation(type="list", formula1=formula,
                                            allow_blank=True)
                worksheet.add_data_validation(validation)
                validation.add(
                    f"{get_column_letter(min_column)}{header_row + 1}:"
                    f"{get_column_letter(max_column)}{last_data_row}"
                )
            return

        clean_column = columns["clean"]
        for item in worksheet.data_validations.dataValidation:
            if item.type == "list" and any(
                area.min_col <= clean_column <= area.max_col
                for area in item.sqref.ranges
            ):
                return
        validation = DataValidation(
            type="list",
            formula1='"Bersih,Kurang Bersih,Tidak Bersih,Not Detected"',
            allow_blank=True,
        )
        worksheet.add_data_validation(validation)
        validation.add(
            f"{get_column_letter(clean_column)}{header_row + 1}:"
            f"{get_column_letter(clean_column)}{last_data_row}"
        )

    @staticmethod
    def _copy_row_style(worksheet: Any, source_row: int, target_row: int) -> None:
        for column in range(1, worksheet.max_column + 1):
            source = worksheet.cell(source_row, column)
            target = worksheet.cell(target_row, column)
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            if source.alignment:
                target.alignment = copy(source.alignment)
        worksheet.row_dimensions[target_row].height = (
            worksheet.row_dimensions[source_row].height or 98
        )

    @staticmethod
    def _write_record(worksheet: Any, row: int, record: ReportItem,
                      columns: dict[str, int], header_row: int,
                      camera_encoding: str, uppercase_status: bool) -> None:
        if columns["no"]:
            worksheet.cell(row, columns["no"]).value = row - header_row
        for key in ("panel_image", "cam1_image", "cam2_image"):
            if columns[key]:
                worksheet.cell(row, columns[key]).value = None
        worksheet.cell(row, columns["site_id"]).value = record.site_id
        worksheet.cell(row, columns["site_name"]).value = record.site_name
        values = {
            "nvr": record.nvr,
            "cam1": record.cam1,
            "cam2": record.cam2,
            "clean": record.clean,
            "note": record.note,
        }
        for key, value in values.items():
            if key in ("cam1", "cam2") and camera_encoding == "ok_nok":
                value = {
                    "Online": "OK",
                    "Offline": "NOK",
                    "Not Detected": "NOT DETECTED",
                    "Online With Note": "OK",
                }.get(value, value)
            elif key == "nvr" and uppercase_status and value:
                value = value.upper()
            column = columns[key]
            if column:
                worksheet.cell(row, column).value = value

    def _set_busy(self, busy: bool) -> None:
        self.busy = busy
        self.main_tabs.setEnabled(not busy)
        if busy:
            self.activity_indicator.set_running()
            self.busy_status_label.setText(self.process_status.text())
            self.busy_progress.setRange(0, 0)
            self.busy_progress.show()
            self.busy_banner.show()
        else:
            self.export_watchdog.stop()
            self.busy_progress.hide()
            self.busy_banner.hide()
            self.activity_indicator.set_idle()
        self.template_button.setEnabled(not busy)
        self.import_workbook_button.setEnabled(not busy)
        self.images_button.setEnabled(not busy)
        self.process_button.setEnabled(not busy)
        self.reprocess_button.setEnabled(not busy)
        self.export_button.setEnabled(not busy)
        self.export_format_combo.setEnabled(not busy)
        self.remove_button.setEnabled(not busy)
        self.paste_button.setEnabled(not busy)
        self.reset_data_button.setEnabled(not busy)
        self.duplicate_button.setEnabled(
            not busy and bool(self._duplicate_item_indices())
        )
        self.site_check_button.setEnabled(
            not busy and bool(self.template_path)
        )

    def closeEvent(self, event: Any) -> None:
        for worker in (self.import_worker, self.worker, self.export_worker):
            if worker is not None and worker.isRunning():
                worker.wait()
        self._clipboard_tempdir.cleanup()
        self._embedded_tempdir.cleanup()
        super().closeEvent(event)


def main() -> int:
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    window = MainWindow()
    window.showMaximized()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
